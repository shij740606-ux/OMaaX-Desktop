"""services/ppt_build.py —— P5 大纲 → .pptx 初稿（python-pptx 本地生成，8 风格版式）。

版式（16:9 横版）：
- title 页      → layouts[0] Title Slide：标题 + 副标题；cover_style.title_align
                  可居中、accent_bar 加色块条
- section 页    → layouts[6] Blank：色块横幅 + 大标题居中（章节分隔页）
- 内容页        → layouts[1]：**加粗** 拆 run、「- 子项」降为 level 1 缩进
- 图占位页      → 左 (figure_split_pct)% bullets + 右侧虚线框
                  （"Fig 1（论文 S1-P12）｜请插入论文中的图"）
- 大数字页      → result 页带 highlight：60pt 数值 + 24pt 副标题 + 底部小 bullets
- reference 页  → layouts[1]，字号 reference_pt
- speaker_notes → notes_slide（evidence 只进备注不进正文）

P5 二次开发 2.1（archetype，借 paper-skills/paper-ppt-skill，MIT）：
- outline 编号 chip 行、「关键词：说明」要点行、2-4 并列短句圆角卡片、
  '→ ' 流程卡+箭头、'> ' 结论强调框、conclusion ≥6 条 50/50 双栏、
  '| ' → coerce table 原生表格；触发全靠 bullets 语法/启发式，零 schema 改动
- 2.3 Tier B：figs/Fig_N.png 真图存在 → frame+contain 插入；否则 Tier A 占位几何不变

风格：data/ppt_styles/{id}.json tokens 是唯一真源（CSS 是前端投影，同步维护）。
fonts.* 是多值字体栈（CSS 语法），pptx 只写首字体 → a:latin + a:ea 双写。
"""
from __future__ import annotations

import json
import logging
import re
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from ..config import get_settings
from .ppt_figure_index import canon_label
from .ppt_outline import CONTENT_TYPES, SCENARIO_LABELS, STYLE_IDS

logger = logging.getLogger(__name__)

_SLIDE_W_IN = 13.333  # 16:9
_SLIDE_H_IN = 7.5

# P5 增强 3：正文标记语法（LLM 输出 → 版面）
_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_SUB_RE = re.compile(r"^\s*(?:[-–—]|·|•)\s+(.*)$")


class BuildError(Exception):
    """PPT 生成失败（VALIDATION | STYLE_NOT_FOUND | BUILD_FAILED）。"""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


# ---------- 风格 tokens ----------

def load_style(style_id: str) -> dict[str, Any]:
    """读 data/ppt_styles/{id}.json 并做最小结构校验。"""
    if style_id not in STYLE_IDS:
        raise BuildError("VALIDATION", f"未知风格 style_id={style_id}")
    path: Path = get_settings().ppt_styles_abs / f"{style_id}.json"
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise BuildError("STYLE_NOT_FOUND", f"风格文件缺失：{path.name}") from exc
    except (OSError, ValueError) as exc:
        raise BuildError("STYLE_NOT_FOUND", f"风格文件不可读：{exc}") from exc
    if not isinstance(data, dict):
        raise BuildError("STYLE_NOT_FOUND", "风格文件必须是 JSON 对象")
    for key in ("palette", "fonts", "sizes"):
        if not isinstance(data.get(key), dict):
            raise BuildError("STYLE_NOT_FOUND", f"风格文件缺少 {key} 段")
    return data


# ---------- 字体/颜色工具 ----------

def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string((hex_color or "#000000").strip().lstrip("#"))


def _first_font(stack: str) -> str:
    """补充 2：多值字体栈 → pptx 只写首字体（缺字体时 Office 主题回退）。"""
    return next((s.strip() for s in (stack or "").split(",") if s.strip()), "Calibri")


def _apply_font(
    run: Any, stack: str, size_pt: float, color_hex: str, *, bold: bool = False,
) -> None:
    name = _first_font(stack)
    font = run.font
    font.name = name            # a:latin
    font.size = Pt(size_pt)
    font.bold = bold
    font.color.rgb = _rgb(color_hex)
    r_pr = font._rPr
    ea = r_pr.find(qn("a:ea"))
    if ea is None:
        ea = r_pr.makeelement(qn("a:ea"), {})   # CJK 走同一字族
        r_pr.append(ea)
    ea.set("typeface", name)


def _set_bg(slide: Any, hex_color: str) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(hex_color)


def _fill_paragraph(
    p: Any, text: str, stack: str, size_pt: float, color_hex: str, *, bold: bool = False,
) -> None:
    """把 **强调** 拆成多 run（强调部分 bold=True）；无标记则单 run 写入。"""
    pos = 0
    wrote = False
    for m in _BOLD_RE.finditer(text):
        if m.start() > pos:
            run = p.add_run()
            run.text = text[pos:m.start()]
            _apply_font(run, stack, size_pt, color_hex, bold=bold)
        run = p.add_run()
        run.text = m.group(1)
        _apply_font(run, stack, size_pt, color_hex, bold=True)
        wrote = True
        pos = m.end()
    if pos < len(text) or not wrote:
        run = p.add_run()
        run.text = text[pos:]
        _apply_font(run, stack, size_pt, color_hex, bold=bold)


def _set_single_text(
    shape: Any, text: str, stack: str, size_pt: float, color_hex: str, *,
    bold: bool = False, align: PP_ALIGN | None = None,
) -> None:
    tf = shape.text_frame
    tf.word_wrap = True
    tf.clear()
    p = tf.paragraphs[0]
    if align is not None:
        p.alignment = align
    _fill_paragraph(p, text, stack, size_pt, color_hex, bold=bold)


# ---------- P5 二次开发 2.1：archetype 底层 primitives ----------
# 几何与触发语法借鉴 paper-skills style-presets / paper-ppt-skill Step7（均 MIT），
# 以 bullets 标记与启发式触发，零 schema 改动；无真图时全部回落 Tier A 文字版式。

_CONTENT_X, _CONTENT_TOP, _CONTENT_BOTTOM, _CONTENT_W = 0.5, 1.7, 6.6, 12.33
_CARD_LIGHT, _CARD_DARK = "#F2F5F9", "#151E2E"
_KV_RE = re.compile(r"^\s*([^：:]{1,10})[：:]\s*(\S.*)$")
_FLOW_PREFIXES = ("→", "->", "➔", "➜")


def _lum(hex_color: str) -> float:
    h = (hex_color or "#ffffff").lstrip("#")
    try:
        r, g, b = (int(h[i:i + 2], 16) for i in (0, 2, 4))
    except ValueError:
        return 1.0
    return (0.2126 * r + 0.7152 * g + 0.0722 * b) / 255


def _strip_md(text: str) -> str:
    return _BOLD_RE.sub(r"\1", text)


def _card_bg(pal: dict[str, Any]) -> str:
    return str(pal["card_bg"]) if pal.get("card_bg") else (
        _CARD_DARK if _lum(pal.get("bg", "#FFFFFF")) < 0.5 else _CARD_LIGHT
    )


def _mm(v: float) -> float:
    """毫米 → 英寸（本文件几何统一英寸；图卡内边距/图注间距按 mm 规范给）。"""
    return v / 25.4


def _px(v: float) -> float:
    """CSS px → 英寸（16:9 页 13.333in ↔ 1280px → 96dpi；圆角 8px 换算用）。"""
    return v / 96.0


def _mix(color: str, other: str, ratio: float) -> str:
    """把 color 按 ratio 向 other 混合（surface 层次推导）。"""
    a = (color or "#ffffff").lstrip("#")
    b = (other or "#ffffff").lstrip("#")
    try:
        ca = [int(a[i:i + 2], 16) for i in (0, 2, 4)]
        cb = [int(b[i:i + 2], 16) for i in (0, 2, 4)]
    except ValueError:
        return color
    return "#" + "".join(f"{int(round(x + (y - x) * ratio)):02X}" for x, y in zip(ca, cb))


def _surface2(pal: dict[str, Any]) -> str:
    """图卡容器底色 = tokens.surface_2（比页面底浅一档）。

    风格 JSON 目前未写 surface_2 → 按 bg 推导：浅底压暗一档、深底提亮一档，
    白底论文原图贴上去不再与页面底色同色（改进 3 的「融合」前提）。
    """
    for key in ("surface_2", "surface-2"):
        if pal.get(key):
            return str(pal[key])
    bg = str(pal.get("bg") or "#FFFFFF")
    return _mix(bg, "#000000", 0.06) if _lum(bg) >= 0.5 else _mix(bg, "#FFFFFF", 0.12)


def _legible(bg: str, *candidates: str | None) -> str:
    """在 bg 上按顺序挑第一个对比足够的颜色（|Δ亮度| ≥0.30 ≈ 大字号 3:1）。

    tokens.muted / section_text 在部分风格里与底色太接近（如深蓝分隔底 + 灰字、
    浅蓝分隔底 + 白字）会直接看不清 → 逐个回退，最后按底色明暗给黑白。
    """
    bg_l = _lum(bg)
    for c in candidates:
        if c and abs(_lum(c) - bg_l) >= 0.30:
            return str(c)
    return "#FFFFFF" if bg_l < 0.5 else "#0F172A"


def _add_rect(
    slide: Any, x: float, y: float, w: float, h: float, *,
    fill: str | None = None, line: str | None = None, line_pt: float = 1.0,
    dash: bool = False, rounded: bool = False,
) -> Any:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(fill)
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = _rgb(line)
        shape.line.width = Pt(line_pt)
        if dash:
            try:
                shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            except (ValueError, TypeError):  # pragma: no cover - 枚举兼容兜底
                logger.debug("dash_style 不可用，退回实线边框")
    else:
        shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    tf.clear()
    return shape


def _add_rule(slide: Any, x: float, y: float, w: float, color: str) -> None:
    _add_rect(slide, x, y, w, 0.012, fill=color)


def _set_bullet_list(
    shape: Any, bullets: list[str], style: dict[str, Any], *,
    size_pt: float | None = None, spacious: bool = False,
) -> None:
    """要点列表渲染：**强调** 拆 bold run；「- 子项」降级 level 1 缩进。
    v3 spacious：默认页加大字号 + 行间距（仅无特殊版式的普通 bullets 页）。"""
    size = size_pt if size_pt is not None else float(style["sizes"].get("bullet_pt", 18))
    if spacious:
        size = max(min(size + 2, 24), 20)
    if len(bullets) > 6:
        size = max(size - 2, 12)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, raw in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(12 if spacious else 6)
        m = _SUB_RE.match(raw)
        if m:
            p.level = 1          # 子 bullet：PowerPoint 原生二级缩进
            raw = m.group(1)
        _fill_paragraph(
            p, raw, style["fonts"].get("body", ""), size,
            style["palette"].get("text", "#000000"),
        )


# —— 触发启发式（纯函数，供 _render_page 分支决策）——

def _takeaway_index(bullets: list[str]) -> int | None:
    """首字符 '> ' 的 bullet → 强调引用框（该条不再进正文列表）。"""
    for i, b in enumerate(bullets):
        if b.strip().startswith("> "):
            return i
    return None


def _sans_takeaway(bullets: list[str]) -> list[str]:
    """hl/fig/table 分支借道 bullets 时剥 '> ' 前缀（框留给 takeaway 分支，
    避免图/大数字页正文出现字面 '> '）。"""
    return [b.strip().lstrip(">").strip() if b.strip().startswith("> ") else b
            for b in bullets]


# —— v3：行内合并 bullet 拆分（LLM 常把 4 条挤成一行「A- B- C- D」）——

_SPLIT_BULLET_RE = re.compile(r"-\s+")


def _split_bullets(bullets: list[str]) -> list[str]:
    """'- A - B - C' / 'A- B- C' → 独立条目；单段 '- 子项' 保持原样
    （交给 _SUB_RE 降级 level 1）。'3-5' 连字符无空白不拆。"""
    out: list[str] = []
    for raw in bullets:
        parts = [p.strip() for p in _SPLIT_BULLET_RE.split(str(raw).strip()) if p.strip()]
        if len(parts) >= 2:
            out.extend(parts)
        else:
            out.append(str(raw))
    return out


# —— v3：按内容形态的 archetype 触发（降低门槛，非等显式标记）——

_HL_NUM_RE = re.compile(r"(?<![\d.,])((?:\d{1,3},)+\d{3}|\d+(?:\.\d+)?%)(?![\d.,])")
_CMP_RE = re.compile(
    r"对比|\bvs\.?\b|优于|高于|低于|大于|小于|上调|下调|↑|↓", re.IGNORECASE,
)   # v3 注：不用「比较」——"比较长/比较多" 等陈述性用法误伤率高


def _auto_highlight(bullets: list[str]) -> tuple[str, str] | None:
    """恰好 1 条含千分位/百分比大数字 → 自动大数字页（value, label）。"""
    hits = [(i, b) for i, b in enumerate(bullets) if _HL_NUM_RE.search(b)]
    if len(hits) != 1:
        return None
    _, b = hits[0]
    m = _HL_NUM_RE.search(b)
    value = m.group(1)
    label = (b[:m.start(1)] + " " + b[m.end(1):]).strip(" ：:，,。 ")
    label = re.sub(r"\s+", " ", label).strip()
    if len(label) < 2:
        label = b.strip() if b.strip() != value else "关键数字"
    return value, label


def _compare_texts(bullets: list[str]) -> list[str] | None:
    """≥2 条含对比语义（vs/对比/上下调…）→ 对照双栏；≤2 条页 1 命中也触发。"""
    if not (2 <= len(bullets) <= 6):
        return None
    hits = sum(1 for b in bullets if _CMP_RE.search(b))
    if hits >= 2 or (hits == 1 and len(bullets) <= 2):
        return [_strip_md(b) for b in bullets]
    return None


def _flow_texts(bullets: list[str]) -> list[str] | None:
    """≥3 条且全部以 '→' 开头 → 流程卡片行；否则 None。"""
    if not (3 <= len(bullets) <= 6):
        return None
    texts: list[str] = []
    for b in bullets:
        t = b.strip()
        for pfx in _FLOW_PREFIXES:
            if t.startswith(pfx):
                t = t[len(pfx):].strip()
                break
        else:
            return None
        if not t:
            return None
        texts.append(t)
    return texts


def _kv_rows(bullets: list[str], ptype: str) -> list[tuple[str, str]] | None:
    """背景/方法页 ≥3 条全部 '关键词：说明' → 要点行（label 加粗 + rule）。"""
    if ptype not in ("background", "method") or not (3 <= len(bullets) <= 6):
        return None
    rows: list[tuple[str, str]] = []
    for b in bullets:
        m = _KV_RE.match(_strip_md(b))
        if not m:
            return None
        rows.append((m.group(1).strip(), m.group(2).strip()))
    return rows


def _card_texts(bullets: list[str], ptype: str) -> list[str] | None:
    """2-4 条并列短句 → 卡片网格（v3 放宽：长度 ≤26 字、允许 ** 强调，
    页型仍限 method/result/conclusion——背景/参考页保持叙述列表）。"""
    if ptype not in ("method", "result", "conclusion") or not (2 <= len(bullets) <= 4):
        return None
    out: list[str] = []
    for b in bullets:
        if _SUB_RE.match(b):
            return None
        t = b.strip()
        if len(t) > 26:
            return None
        out.append(t)
    return out


def _figure_image(context: dict[str, Any] | None, label: str) -> Path | None:
    """figs/Fig_1.png 真图查找（extract_figures 落盘约定；不存在 → None）。

    v3 修复2 兜底：字符串精确匹配失败（如 "Fig. 6b" 无 Fig_6b.png）→
    按 canon_label（kind+编号）在 figs 目录里找同号文件（Fig 6b → Fig_6.png）。
    """
    figs_dir = (context or {}).get("figs_dir")
    if not figs_dir or not label.strip():
        return None
    path = Path(figs_dir) / f"{label.replace(' ', '_')}.png"
    if path.is_file():
        return path
    key = canon_label(label)
    if not key or not Path(figs_dir).is_dir():
        return None
    for cand in sorted(Path(figs_dir).glob("*.png")):
        if canon_label(cand.stem.replace("_", " ")) == key:
            return cand
    return None


def _first_real_figure(context: dict[str, Any] | None, page: dict[str, Any]) -> Path | None:
    for f in page.get("suggested_figures") or []:
        if isinstance(f, dict):
            hit = _figure_image(context, str(f.get("label") or ""))
            if hit:
                return hit
    return None


# —— archetype 渲染器 ——

def _add_outline_chips(slide: Any, bullets: list[str], style: dict[str, Any]) -> None:
    """outline 页：编号 chip + 结论式条目（借 paper-skills outline archetype）。"""
    pal, fonts = style["palette"], style["fonts"]
    chip_fg = str(pal.get("bg") or "#FFFFFF")
    texts = [_strip_md(str(b)) for b in bullets][:8]
    n = max(len(texts), 1)
    step = min(0.85, (_CONTENT_BOTTOM - _CONTENT_TOP - 0.4) / n)
    y = _CONTENT_TOP + 0.1
    for i, text in enumerate(texts):
        chip = _add_rect(slide, 0.9, y, 0.5, 0.5, fill=pal.get("primary", "#000000"), rounded=True)
        tf = chip.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = str(i + 1)
        _apply_font(run, fonts.get("title", ""), 16, chip_fg, bold=True)
        tb = slide.shapes.add_textbox(Inches(1.7), Inches(y), Inches(10.6), Inches(0.5))
        _set_single_text(tb, text, fonts.get("body", ""), 20, pal.get("text", "#000000"))
        y += step


def _add_kv_rows(slide: Any, rows: list[tuple[str, str]], style: dict[str, Any]) -> None:
    pal, fonts = style["palette"], style["fonts"]
    n = max(len(rows), 1)
    step = min(1.0, (_CONTENT_BOTTOM - _CONTENT_TOP) / n)
    y = _CONTENT_TOP
    for label, value in rows:
        lb = slide.shapes.add_textbox(Inches(0.9), Inches(y + 0.08), Inches(2.5), Inches(step - 0.12))
        _set_single_text(lb, label, fonts.get("body", ""), 18, pal.get("primary", "#000000"), bold=True)
        vb = slide.shapes.add_textbox(Inches(3.6), Inches(y + 0.08), Inches(8.7), Inches(step - 0.12))
        _set_single_text(vb, value, fonts.get("body", ""), 17, pal.get("text", "#000000"))
        if y + step < _CONTENT_BOTTOM - 0.05:
            _add_rule(slide, 0.9, y + step - 0.02, 11.4, pal.get("box_line", "#999999"))
        y += step


def _add_cards(slide: Any, texts: list[str], style: dict[str, Any]) -> None:
    """2-4 张圆角卡片（1×2 / 1×3 / 2×2），几何借 paper-skills key-message 卡。"""
    pal, fonts = style["palette"], style["fonts"]
    n = len(texts)
    cols = n if n <= 3 else 2
    rows = 1 if n <= 3 else 2
    gap = 0.3
    cw = (_CONTENT_W - gap * (cols - 1)) / cols
    ch = (_CONTENT_BOTTOM - _CONTENT_TOP - gap * (rows - 1)) / rows
    for idx, text in enumerate(texts):
        r, c = divmod(idx, cols)
        x = _CONTENT_X + c * (cw + gap)
        y = _CONTENT_TOP + r * (ch + gap)
        card = _add_rect(
            slide, x, y, cw, ch, fill=_card_bg(pal),
            line=pal.get("primary", "#999999"), line_pt=1.25, rounded=True,
        )
        tf = card.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        _fill_paragraph(p, text, fonts.get("body", ""), 17, pal.get("text", "#000000"))


def _add_flow(slide: Any, texts: list[str], style: dict[str, Any]) -> None:
    """流程步骤卡 + 箭头（借 paper-skills buildSteps archetype）。"""
    pal, fonts = style["palette"], style["fonts"]
    n = len(texts)
    arrow_w = 0.55
    cw = (_CONTENT_W - arrow_w * (n - 1)) / n
    ch, y = 2.6, _CONTENT_TOP + 1.1
    x = _CONTENT_X
    for i, text in enumerate(texts):
        card = _add_rect(
            slide, x, y, cw, ch, fill=_card_bg(pal),
            line=pal.get("primary", "#999999"), line_pt=1.25, rounded=True,
        )
        tf = card.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        _fill_paragraph(p, text, fonts.get("body", ""), 16, pal.get("text", "#000000"))
        x += cw
        if i < n - 1:
            ar = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y + ch / 2 - 0.22),
                Inches(arrow_w), Inches(0.44),
            )
            ar.fill.solid()
            ar.fill.fore_color.rgb = _rgb(pal.get("accent", "#666666"))
            ar.line.fill.background()
            x += arrow_w


def _add_takeaway_box(slide: Any, text: str, style: dict[str, Any], *, top: float, height: float) -> None:
    """强调引用框（'> ' 结论句）：accent 边框 + 卡底。"""
    pal, fonts = style["palette"], style["fonts"]
    box = _add_rect(
        slide, 0.9, top, 11.53, height, fill=_card_bg(pal),
        line=pal.get("accent", "#666666"), line_pt=2.25, rounded=True,
    )
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _fill_paragraph(p, _strip_md(text), fonts.get("body", ""), 18, pal.get("text", "#000000"), bold=True)


def _add_two_col(slide: Any, left_bullets: list[str], right_bullets: list[str],
                 style: dict[str, Any], body: Any) -> None:
    """conclusion ≥6 条 → 50/50 双栏（借 paper-ppt-skill Step7 左右均分）。"""
    body.text_frame.clear()
    tb = slide.shapes.add_textbox(Inches(_CONTENT_X), Inches(_CONTENT_TOP), Inches(5.9), Inches(4.9))
    _set_bullet_list(tb, left_bullets, style)
    tb2 = slide.shapes.add_textbox(Inches(6.93), Inches(_CONTENT_TOP), Inches(5.9), Inches(4.9))
    _set_bullet_list(tb2, right_bullets, style)


def _add_compare(slide: Any, texts: list[str], style: dict[str, Any]) -> None:
    """v3 对照布局：两列卡片 + 中缝 accent 竖线（2 条时中央 VS 圆）。"""
    pal, fonts = style["palette"], style["fonts"]
    n = len(texts)
    mid = (n + 1) // 2
    cols = (texts[:mid], texts[mid:])
    gap = 0.7
    col_w = (_CONTENT_W - gap) / 2
    gap_v = 0.24
    rows = max(len(c) for c in cols) or 1
    ch = (_CONTENT_BOTTOM - _CONTENT_TOP - gap_v * (rows - 1)) / rows
    for ci, col in enumerate(cols):
        x = _CONTENT_X + ci * (col_w + gap)
        for ri, text in enumerate(col):
            y = _CONTENT_TOP + ri * (ch + gap_v)
            card = _add_rect(
                slide, x, y, col_w, ch, fill=_card_bg(pal),
                line=pal.get("primary", "#999999"), line_pt=1.25, rounded=True,
            )
            tf = card.text_frame
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = Inches(0.18)
            tf.margin_right = Inches(0.18)
            p = tf.paragraphs[0]
            _fill_paragraph(p, text, fonts.get("body", ""), 17, pal.get("text", "#000000"))
    mid_x = _CONTENT_X + col_w + gap / 2
    _add_rect(slide, mid_x - 0.015, _CONTENT_TOP, 0.03,
              _CONTENT_BOTTOM - _CONTENT_TOP, fill=pal.get("accent", "#2563EB"))
    if n == 2:
        d = 0.74
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(mid_x - d / 2),
            Inches((_CONTENT_TOP + _CONTENT_BOTTOM) / 2 - d / 2),
            Inches(d), Inches(d),
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = _rgb(pal.get("accent", "#2563EB"))
        circle.line.color.rgb = _rgb(pal.get("bg", "#FFFFFF"))
        circle.line.width = Pt(3.0)
        ctf = circle.text_frame
        ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
        cp = ctf.paragraphs[0]
        cp.alignment = PP_ALIGN.CENTER
        crun = cp.add_run()
        crun.text = "VS"
        _apply_font(crun, fonts.get("title", ""), 18, "#FFFFFF", bold=True)


def _valid_table(page: dict[str, Any]) -> bool:
    t = page.get("table")
    return isinstance(t, dict) and isinstance(t.get("rows"), list) and bool(t["rows"])


def _add_data_table(slide: Any, table: dict[str, Any], style: dict[str, Any],
                    *, top: float, height: float) -> None:
    """原生 pptx 表格：primary 表头 + 隔行卡底（'| ' bullet → coerce 成 table）。"""
    pal, fonts = style["palette"], style["fonts"]
    headers = [str(h) for h in (table.get("headers") or [])]
    rows = [[str(c) for c in r] for r in table["rows"]]
    cols = max([len(headers)] + [len(r) for r in rows]) or 1
    headers += [""] * (cols - len(headers))
    rows = [r + [""] * (cols - len(r)) for r in rows]
    shape = slide.shapes.add_table(
        len(rows) + 1, cols, Inches(0.7), Inches(top), Inches(11.9), Inches(height),
    )
    tbl = shape.table
    for c in range(cols):
        tbl.columns[c].width = Inches(11.9 / cols)
    size = 14 if cols <= 5 else 12
    header_fg = "#FFFFFF" if _lum(pal.get("primary", "#000000")) < 0.45 else str(pal.get("bg") or "#FFFFFF")
    bg = _card_bg(pal)

    def _fill_cell(cell: Any, text: str, *, bold: bool, fg: str, back: str,
                   size_pt: float, align: PP_ALIGN) -> None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = _rgb(back)
        tf = cell.text_frame
        tf.word_wrap = True
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = align
        run = p.add_run()
        run.text = text
        _apply_font(run, fonts.get("body", ""), size_pt, fg, bold=bold)

    for c, htext in enumerate(headers):
        _fill_cell(tbl.cell(0, c), htext, bold=True, fg=header_fg,
                   back=pal.get("primary", "#000000"), size_pt=size, align=PP_ALIGN.CENTER)
    for r, row in enumerate(rows):
        back = bg if r % 2 == 0 else str(pal.get("bg") or "#FFFFFF")
        for c, text in enumerate(row):
            _fill_cell(tbl.cell(r + 1, c), text, bold=False, fg=pal.get("text", "#000000"),
                       back=back, size_pt=size, align=PP_ALIGN.LEFT)


def _add_real_figure(slide: Any, img_path: Path, figs: list[dict], style: dict[str, Any],
                     *, left_in: float, top_in: float, width_in: float, height_in: float) -> None:
    """Tier B：真图 contain 插入 —— 浅色圆角卡片容器 + 图 + 底部图注（改进 3）。

    白底论文原图不再直接贴页面底色：容器 tokens.surface_2 + 圆角 8px + 无边框，
    图四周 5mm 内边距并居中；图注 10-12pt tokens.muted、贴容器底部、与图间距 3mm，
    格式「Fig 3（论文 S2-P19）」。
    """
    pal, fonts = style["palette"], style["fonts"]
    fig_pt = min(max(float(style["sizes"].get("figure_pt", 13)), 10.0), 12.0)
    # 1) 圆角矩形容器（surface_2，无边框）
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left_in), Inches(top_in), Inches(width_in), Inches(height_in),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = _rgb(_surface2(pal))
    card.line.fill.background()
    try:  # 圆角半径 8px（96dpi → 0.083in）；默认 adj 会圆成胶囊，按短边换算
        card.adjustments[0] = min(0.5, _px(8) / min(width_in, height_in))
    except (IndexError, ValueError, TypeError):  # pragma: no cover - 调整点兼容兜底
        logger.debug("圆角调整不可用，保留默认圆角")
    tf = card.text_frame
    tf.word_wrap = True
    tf.clear()

    # 2) 图注（容器底部）：逐图一行，与图区留 3mm
    captions = [
        f"{f.get('label') or '图'}（论文 {f.get('para_id') or '—'}）" for f in figs[:6]
    ]
    cap_h = (len(captions) * 0.2 + _mm(3)) if captions else 0.0
    pad = _mm(5)
    ix, iy = left_in + pad, top_in + pad
    iw = width_in - 2 * pad
    ih = height_in - 2 * pad - cap_h
    aspect: float | None = None
    try:
        from PIL import Image
        with Image.open(img_path) as im:
            if im.height:
                aspect = im.width / im.height
    except Exception:  # noqa: BLE001 - 图片不可读则只留容器
        logger.debug("真图不可读：%s", img_path)
    if aspect and ih > 0.2 and iw > 0.2:
        w, h = (ih * aspect, ih) if iw / ih > aspect else (iw, iw / aspect)
        try:
            slide.shapes.add_picture(
                str(img_path), Inches(ix + (iw - w) / 2), Inches(iy + (ih - h) / 2),
                Inches(w), Inches(h),
            )
        except OSError:  # pragma: no cover
            logger.debug("真图插入失败：%s", img_path)
    if not captions:
        return
    cap_tb = slide.shapes.add_textbox(
        Inches(left_in), Inches(top_in + height_in - cap_h),
        Inches(width_in), Inches(cap_h),
    )
    ctf = cap_tb.text_frame
    ctf.word_wrap = True
    ctf.margin_left = ctf.margin_right = Inches(0.06)
    ctf.margin_top = ctf.margin_bottom = 0
    ctf.clear()
    cap_color = _legible(_surface2(pal), pal.get("muted"), pal.get("text"), pal.get("section_text"))
    for i, text in enumerate(captions):
        p = ctf.paragraphs[0] if i == 0 else ctf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        _apply_font(r, fonts.get("body", ""), fig_pt, cap_color)


def _render_table_page(slide: Any, page: dict[str, Any], style: dict[str, Any],
                       body: Any, bullets: list[str]) -> None:
    """table 页：≤2 条导读 bullets 置顶 + 原生表格吃剩余版面。"""
    top_b = bullets[:2]
    t_top = _CONTENT_TOP
    if top_b:
        body.left, body.top = Inches(0.7), Inches(_CONTENT_TOP)
        body.width = Inches(11.9)
        body.height = Inches(0.5 * len(top_b) + 0.12)
        _set_bullet_list(body, _sans_takeaway(top_b), style,
                         size_pt=max(float(style["sizes"].get("bullet_pt", 18)) - 4, 12))
        t_top = _CONTENT_TOP + 0.5 * len(top_b) + 0.24
    else:
        body.text_frame.clear()
    if t_top > 5.0:
        body.text_frame.clear()
        t_top = _CONTENT_TOP
    _add_data_table(slide, page["table"], style, top=t_top,
                    height=max(_CONTENT_BOTTOM - t_top, 1.0))


# ---------- 页渲染 ----------

def _cover_fallback(context: dict[str, Any] | None) -> str:
    if not context:
        return "科研汇报"
    label = SCENARIO_LABELS.get(str(context.get("scenario") or ""), "科研汇报")
    minutes = context.get("minutes")
    return f"{label} · {minutes} 分钟" if minutes else label


def _notes_text(page: dict[str, Any]) -> str:
    lines: list[str] = []
    if page.get("speaker_notes"):
        lines.append(str(page["speaker_notes"]))
    for e in page.get("evidence") or []:
        if isinstance(e, dict) and e.get("quote"):
            lines.append(
                f"来源 {e.get('para_id') or '—'}：{e['quote']}"
                f"（{e.get('status') or 'unverified'}）"
            )
    for f in page.get("suggested_figures") or []:
        if isinstance(f, dict):
            lines.append(
                f"图 {f.get('label') or '—'} → {f.get('para_id') or '—'}"
                f"（{f.get('status') or 'unverified'}）"
            )
    return "\n".join(lines)


def _add_notes(slide: Any, page: dict[str, Any]) -> None:
    text = _notes_text(page)
    if text:
        slide.notes_slide.notes_text_frame.text = text


def _set_bullets(shape: Any, page: dict[str, Any], style: dict[str, Any],
                 *, size_pt: float | None = None, spacious: bool = False) -> None:
    bullets = _split_bullets([str(b) for b in (page.get("bullets") or [])])
    _set_bullet_list(shape, bullets, style, size_pt=size_pt, spacious=spacious)


def _add_figure_box(slide: Any, page: dict[str, Any], style: dict[str, Any], *,
                    left_in: float, width_in: float,
                    top_in: float = 1.7, height_in: float = 4.9) -> None:
    """右侧虚线图占位框：'Fig 1（论文 S1-P12）'+ 插图提示（手工插图模式）。"""
    pal = style["palette"]
    fig_pt = float(style["sizes"].get("figure_pt", 13))
    box = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left_in), Inches(top_in), Inches(width_in), Inches(height_in),
    )
    box.fill.background()  # 透明，只留边框
    box.line.color.rgb = _rgb(pal.get("box_line", "#999999"))
    box.line.width = Pt(1.5)
    try:
        box.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    except (ValueError, TypeError):  # pragma: no cover - 枚举兼容兜底
        logger.debug("dash_style 不可用，退回实线边框")
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.clear()
    figs = [f for f in (page.get("suggested_figures") or []) if isinstance(f, dict)][:6]
    first = True
    for fig in figs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = f"{fig.get('label') or '图'}（论文 {fig.get('para_id') or '—'}）"
        _apply_font(run, style["fonts"].get("body", ""), fig_pt, pal.get("primary", "#000000"), bold=True)
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        run2 = p2.add_run()
        run2.text = "请插入论文中的图"
        _apply_font(run2, style["fonts"].get("body", ""), max(fig_pt - 2, 9), pal.get("muted", "#666666"))
        p2.space_after = Pt(8)


def _add_highlight(slide: Any, page: dict[str, Any], style: dict[str, Any]) -> None:
    """大数字突出页（借鉴 ppt-agent-skill "big number"）：60pt 数值 + 24pt 副标题。"""
    hl = page.get("highlight") or {}
    pal, fonts = style["palette"], style["fonts"]
    cs = style.get("content_style") or {}
    value_pt = float(cs.get("highlight_value_pt", 60))
    label_pt = float(cs.get("highlight_label_pt", 24))

    value_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.7), Inches(12.33), Inches(1.5))
    _set_single_text(
        value_box, str(hl.get("value") or ""), fonts.get("title", ""), value_pt,
        pal.get("accent", "#000000"), bold=True, align=PP_ALIGN.CENTER,
    )
    label_box = slide.shapes.add_textbox(Inches(0.5), Inches(3.25), Inches(12.33), Inches(0.7))
    _set_single_text(
        label_box, str(hl.get("label") or ""), fonts.get("body", ""), label_pt,
        pal.get("text", "#000000"), align=PP_ALIGN.CENTER,
    )


def _section_subtitle(page: dict[str, Any]) -> str:
    """分隔页副标题：subtitle 字段 → bullets[0]（旧大纲）→ 空。"""
    sub = str(page.get("subtitle") or "").strip()
    if sub:
        return _strip_md(sub)
    return next(
        (_strip_md(str(b).strip()) for b in (page.get("bullets") or []) if str(b).strip()), "",
    )


def _add_section_page(slide: Any, page: dict[str, Any], style: dict[str, Any]) -> None:
    """章节分隔页（P5 改进 2，layouts[6] Blank）：
    section_bg 底色 + 40pt 主标题居中粗体 + 20pt 副标题（tokens.muted）
    + 底部 30% 宽 accent 装饰线。
    """
    pal, fonts, sizes = style["palette"], style["fonts"], style["sizes"]
    section_bg = str(
        pal.get("section_bg")
        or pal.get("surface_2") or pal.get("surface-2")
        or pal.get("cover_bg") or pal.get("bg") or "#FFFFFF"
    )
    _set_bg(slide, section_bg)

    title_color = _legible(
        section_bg, pal.get("text"), pal.get("cover_text"), pal.get("section_text"),
    )
    title_box = slide.shapes.add_textbox(
        Inches(0.6), Inches(2.85), Inches(_SLIDE_W_IN - 1.2), Inches(1.1),
    )
    ttf = title_box.text_frame
    ttf.word_wrap = True
    ttf.vertical_anchor = MSO_ANCHOR.MIDDLE
    ttf.clear()
    tp = ttf.paragraphs[0]
    tp.alignment = PP_ALIGN.CENTER
    trun = tp.add_run()
    trun.text = str(page.get("title") or "章节")
    _apply_font(
        trun, fonts.get("title", ""), float(sizes.get("section_title_pt", 40)),
        title_color, bold=True,
    )

    subtitle = _section_subtitle(page)
    if subtitle:
        sub_color = _legible(
            section_bg, pal.get("muted"), pal.get("section_text"), pal.get("text"),
        )
        sub_box = slide.shapes.add_textbox(
            Inches(0.9), Inches(4.05), Inches(_SLIDE_W_IN - 1.8), Inches(0.7),
        )
        stf = sub_box.text_frame
        stf.word_wrap = True
        stf.vertical_anchor = MSO_ANCHOR.TOP
        stf.clear()
        sp = stf.paragraphs[0]
        sp.alignment = PP_ALIGN.CENTER
        srun = sp.add_run()
        srun.text = subtitle
        _apply_font(srun, fonts.get("body", ""), 20.0, sub_color)

    # 底部装饰线：accent 色、宽 30% 居中
    line_w = _SLIDE_W_IN * 0.3
    _add_rect(
        slide, (_SLIDE_W_IN - line_w) / 2, 6.62, line_w, 0.03,
        fill=_legible(section_bg, pal.get("accent"), pal.get("primary")),
    )


def _add_cover_bar(slide: Any, style: dict[str, Any], *, centered: bool) -> None:
    """cover_style.accent_bar：居中封面 → 底部通栏；左对齐封面 → 左侧竖条。"""
    pal = style["palette"]
    if centered:
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(_SLIDE_H_IN - 0.15), Inches(_SLIDE_W_IN), Inches(0.15),
        )
    else:
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.18), Inches(_SLIDE_H_IN),
        )
    bar.fill.solid()
    bar.fill.fore_color.rgb = _rgb(pal.get("accent", "#2563EB"))
    bar.line.fill.background()


def _render_page(
    prs: Presentation, page: dict[str, Any], style: dict[str, Any],
    context: dict[str, Any] | None,
) -> str:
    """渲染单页，返回 archetype 标签（v3：供 stats.archetypes 覆盖率统计）。"""
    pal, fonts, sizes = style["palette"], style["fonts"], style["sizes"]
    ptype = page.get("type")

    if ptype == "title":
        cover_style = style.get("cover_style") or {}
        centered = str(cover_style.get("title_align") or "left") == "center"
        align = PP_ALIGN.CENTER if centered else None
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        _set_bg(slide, pal.get("cover_bg", "#FFFFFF"))
        _set_single_text(
            slide.shapes.title, str(page.get("title") or "科研汇报"),
            fonts.get("title", ""), float(sizes.get("cover_title_pt", 40)),
            pal.get("cover_text", "#000000"), bold=True, align=align,
        )
        subtitle = " ｜ ".join(str(b) for b in (page.get("bullets") or [])) or _cover_fallback(context)
        _set_single_text(
            slide.placeholders[1], subtitle,
            fonts.get("body", ""), float(sizes.get("bullet_pt", 18)),
            pal.get("muted", "#666666"), align=align,
        )
        if cover_style.get("accent_bar"):
            _add_cover_bar(slide, style, centered=centered)
        _add_notes(slide, page)
        return "cover"

    if ptype == "section":
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
        _add_section_page(slide, page, style)
        _add_notes(slide, page)
        return "section"

    has_fig = bool(page.get("suggested_figures"))
    has_hl = bool(page.get("highlight"))
    # v3：行内合并 bullet 拆分（'A- B- C' → 独立条目）
    bullets = _split_bullets([str(b) for b in (page.get("bullets") or [])])
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    _set_bg(slide, pal.get("bg", "#FFFFFF"))
    _set_single_text(
        slide.shapes.title, str(page.get("title") or ptype),
        fonts.get("title", ""), float(sizes.get("title_pt", 30)),
        pal.get("primary", "#000000"), bold=True,
    )
    body = slide.placeholders[1]
    if ptype == "outline" and bullets:
        # 2.1：outline archetype（编号 chip 行，paper-skills outline 版式）
        body.text_frame.clear()
        _add_outline_chips(slide, bullets, style)
        _add_notes(slide, page)
        return "chips"
    if has_hl:
        # 大数字页：数值/副标题居中置顶，bullets 缩到下部小字
        _add_highlight(slide, page, style)
        body.left, body.top = Inches(0.9), Inches(4.35)
        body.width, body.height = Inches(11.5), Inches(2.5)
        if bullets:
            _set_bullet_list(body, _sans_takeaway(bullets), style,
                             size_pt=max(float(sizes.get("bullet_pt", 18)) - 4, 12))
        else:
            body.text_frame.clear()
        _add_notes(slide, page)
        return "highlight"
    if _valid_table(page):
        # 2.1：对比表格（'| ' bullets → coerce 成 table）；优先于图占位——
        # 同页既引用图又给表时，表是显式数据内容，图证据仍进 notes
        _render_table_page(slide, page, style, body, bullets)
        _add_notes(slide, page)
        return "table"
    if has_fig:
        # 图文混排：左 figure_split_pct% bullets + 右侧图卡（改进 3：圆角卡片容器）
        # Tier B：figs/ 下有真图 → 容器内 contain 插入；否则 Tier A 虚线占位
        real = _first_real_figure(context, page)
        figs = [f for f in page.get("suggested_figures") or [] if isinstance(f, dict)][:6]
        if bullets:
            usable, gap = 12.33, 0.3
            cs = style.get("content_style") or {}
            split = min(max(float(cs.get("figure_split_pct", 60)), 45), 75) / 100
            body_w = (usable - gap) * split
            box_left = 0.5 + body_w + gap
            # 图占位控制：右侧 ≤55% 页宽、≤60% 页高
            box_w = min(12.83 - box_left, _SLIDE_W_IN * 0.55)
            box_top, box_h = 1.7, _SLIDE_H_IN * 0.6
            body.left, body.top = Inches(0.5), Inches(1.7)
            body.width, body.height = Inches(body_w), Inches(4.9)
            _set_bullet_list(body, _sans_takeaway(bullets), style,
                             size_pt=float(sizes.get("bullet_pt", 18)) - 2)
        else:
            # 单独图页：图居中，上下留白 ≥5% 页高
            body.text_frame.clear()
            box_w = _SLIDE_W_IN * 0.65
            box_left = (_SLIDE_W_IN - box_w) / 2
            box_top, box_h = 1.9, _SLIDE_H_IN * 0.6
        if real is not None:
            _add_real_figure(slide, real, figs, style,
                             left_in=box_left, top_in=box_top,
                             width_in=box_w, height_in=box_h)
        else:
            _add_figure_box(slide, page, style, left_in=box_left, width_in=box_w,
                            top_in=box_top, height_in=box_h)
        _add_notes(slide, page)
        return "figure"
    if (ti := _takeaway_index(bullets)) is not None:
        # 2.1：'> ' 结论句 → 底部强调引用框，其余 bullets 留正文
        rest = [b for i, b in enumerate(bullets) if i != ti]
        box_h, gap_in = 1.0, 0.18
        box_top = _CONTENT_BOTTOM - box_h
        _add_takeaway_box(slide, bullets[ti].lstrip(">").strip(), style,
                          top=box_top, height=box_h)
        body.left, body.top = Inches(0.9), Inches(_CONTENT_TOP)
        body.width, body.height = Inches(11.53), Inches(max(box_top - gap_in - _CONTENT_TOP, 0.4))
        if rest:
            _set_bullet_list(body, rest, style,
                             size_pt=max(float(sizes.get("bullet_pt", 18)) - 2, 12))
        else:
            body.text_frame.clear()
        _add_notes(slide, page)
        return "takeaway"
    if (flow := _flow_texts(bullets)) is not None:
        # 2.1：'→ ' 步骤 → 流程卡 + 箭头
        body.text_frame.clear()
        _add_flow(slide, flow, style)
        _add_notes(slide, page)
        return "flow"
    if (kv := _kv_rows(bullets, ptype)) is not None:
        # 2.1：'关键词：说明' → label 加粗要点行
        body.text_frame.clear()
        _add_kv_rows(slide, kv, style)
        _add_notes(slide, page)
        return "kv"
    if (cmp_texts := _compare_texts(bullets)) is not None:
        # v3：对比语义（vs/对比/上下调…）→ 对照双栏
        body.text_frame.clear()
        _add_compare(slide, cmp_texts, style)
        _add_notes(slide, page)
        return "compare"
    if (auto_hl := _auto_highlight(bullets)) is not None:
        # v3：恰 1 条含千分位/百分比大数字 → 自动大数字页（LLM 未显式给 highlight）
        _add_highlight(slide, dict(page, highlight={"value": auto_hl[0], "label": auto_hl[1]}),
                       style)
        body.left, body.top = Inches(0.9), Inches(4.35)
        body.width, body.height = Inches(11.5), Inches(2.5)
        if bullets:
            _set_bullet_list(body, _sans_takeaway(bullets), style,
                             size_pt=max(float(sizes.get("bullet_pt", 18)) - 4, 12))
        else:
            body.text_frame.clear()
        _add_notes(slide, page)
        return "highlight_auto"
    if (cards := _card_texts(bullets, ptype)) is not None:
        # 2.1：并列短句 → 圆角卡片网格（v3 放宽页型/长度/允许 **）
        body.text_frame.clear()
        _add_cards(slide, cards, style)
        _add_notes(slide, page)
        return "cards"
    if ptype == "conclusion" and len(bullets) >= 6:
        # 2.1：结论 ≥6 条 → 50/50 双栏
        mid = (len(bullets) + 1) // 2
        _add_two_col(slide, bullets[:mid], bullets[mid:], style, body)
        _add_notes(slide, page)
        return "two_col"
    size = float(sizes.get("reference_pt", 14)) if ptype == "reference" else None
    # v3：普通 bullets 页加大字号 + 行间距（reference 保持小字）
    _set_bullets(body, page, style, size_pt=size, spacious=ptype != "reference")
    _add_notes(slide, page)
    return "reference" if ptype == "reference" else "bullets"


# ---------- 主入口 ----------

def section_subtitles(pages: list[dict[str, Any]]) -> dict[int, str]:
    """section 页副标题决议（改进 2）：subtitle → bullets[0] → 该章第一页 title 兜底。

    「该章」= 该分隔页之后到下一张分隔页之前的第一张内容页（background/method/
    result/conclusion）；返回 {页下标: 副标题}（空串表示该页无副标题）。
    """
    out: dict[int, str] = {}
    for i, page in enumerate(pages):
        if page.get("type") != "section":
            continue
        sub = _section_subtitle(page)
        if not sub:
            for nxt in pages[i + 1:]:
                if nxt.get("type") == "section":
                    break
                if nxt.get("type") in CONTENT_TYPES:
                    title = str(nxt.get("title") or "").strip()
                    if title:
                        sub = _strip_md(title)
                        break
        out[i] = sub
    return out


def list_styles() -> list[dict[str, Any]]:
    """风格清单（data/ppt_styles/*.json 实际内容；改进 1：前端数量不再写死）。

    单个风格文件缺失/损坏 → 跳过该条（清单其余照常返回，端点不 500）。
    """
    items: list[dict[str, Any]] = []
    for sid in STYLE_IDS:
        try:
            style = load_style(sid)
        except BuildError as exc:
            logger.warning("风格清单跳过 %s：%s", sid, exc.message)
            continue
        pal = style.get("palette") or {}
        items.append({
            "id": sid,
            "label": str(style.get("label") or style.get("display_name") or sid),
            "display_name": str(style.get("display_name") or sid),
            "description": str(style.get("description") or ""),
            "bg": str(pal.get("bg") or "#FFFFFF"),
            "accent": str(pal.get("accent") or "#2563EB"),
        })
    return items


def build_pptx(
    pages: list[dict], style_id: str, out_path: Path,
    *, context: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """大纲 → .pptx 落盘。返回 {"slides": n, "bytes": n}；失败抛 BuildError。"""
    if not pages:
        raise BuildError("VALIDATION", "大纲为空，无法生成 PPT")
    style = load_style(style_id)
    prs = Presentation()
    prs.slide_width = Inches(_SLIDE_W_IN)
    prs.slide_height = Inches(_SLIDE_H_IN)
    # 2.3：真图目录 = 输出同级 figs/（extract_figures 落盘约定；缺目录 → Tier A 占位）
    ctx = dict(context or {})
    ctx.setdefault("figs_dir", out_path.parent / "figs")
    subs = section_subtitles(pages)
    archetypes: list[str] = []
    for i, page in enumerate(pages):
        if page.get("type") == "section":
            page = {**page, "subtitle": subs.get(i, "")}   # 决议后的副标题（含兜底）
        archetypes.append(_render_page(prs, page, style, ctx))
    try:
        out_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(out_path))
    except OSError as exc:
        raise BuildError("BUILD_FAILED", f"PPTX 落盘失败：{exc}") from exc
    stats = {
        "slides": len(prs.slides._sldIdLst),
        "bytes": out_path.stat().st_size,
        # v3：逐页 archetype（"bullets"/"reference" 之外 = 版式变化页）
        "archetypes": archetypes,
    }
    logger.info("pptx 生成：%s slides=%s bytes=%s", out_path.name, stats["slides"], stats["bytes"])
    return stats
