"""services/format_check.py —— 版式与规范规则引擎（⑤ 论文排版与规范）。

纯规则、无 LLM、无 vstore、同步执行（通常 <1s）：
- PDF   → pymupdf 本地解档（AGENTS.md 约束 6 例外：vstore 只回文本，字体/边距/
  行距/页数等版式元数据拿不到；正文提取仍走 vstore，本服务不碰正文语义）；
- Word  → python-docx（页边距/行距精确，页数取 docProps/app.xml 的 <Pages>）；
- 规则集 → data/journals/journals.json：
    顶层 format_rules_default（通用默认值，rules_source="default"）
    每条目可选 format_rules（按期刊覆盖，必须带 rules_source: placeholder|official|default）
  ⚠ 诚实原则：不知道期刊真实投稿指南 → 该条只写 {"rules_source": "placeholder"}，
  所有具体数字回退 default 通用值，绝不编造期刊专属数字。

输出：{checks[{name,key,status,detail,suggestion}], summary{pass,warn,fail}, meta{...}}
status ∈ pass|warn|fail。参考文献样式宽容化：仅"明确不符"（如 [J] 标引但规则要求 APA）
才 fail，"疑似但不确定/检测不到"一律 warn，避免误伤。
"""
from __future__ import annotations

import io
import json
import re
import statistics
import zipfile
from pathlib import Path
from typing import Any

import fitz  # pymupdf（venv 已装 1.28.x）
from docx import Document
from docx.shared import Length

# review_service/app/services/ → parents[2] = review_service
_JOURNALS_PATH = Path(__file__).resolve().parents[2] / "data" / "journals" / "journals.json"

_MAX_SCAN_PAGES = 80   # 版式扫描页数上限（正文/字体/边距；稿件通常远小于此）
_MAX_REFS_SCAN = 40    # 从文末向前找参考文献标题的页数窗口
_MAX_PARAS = 8000      # docx 段落扫描上限
_PT_TO_MM = 25.4 / 72.0
_EMU_TO_MM = 1.0 / 36000.0

_HEADING_RE = re.compile(
    r"^\s*(references|bibliography|literature cited|works cited|参考文献|文献)\s*[:：]?\s*$",
    re.IGNORECASE,
)
_FIG_RE = re.compile(r"(?im)^\s*(figure|fig\.?|图)\s*\d+")
_TAB_RE = re.compile(r"(?im)^\s*(table|表)\s*\d+")
_GB_TAG_RE = re.compile(r"\[[A-Z]{1,3}\d*\]")          # [J] [M] [D] [C] …（GB/T 7714 标引）
_PAREN_YEAR_RE = re.compile(r"\((?:19|20)\d{2}\)")      # (2024) 作者-年份制
_BRACKET_NUM_RE = re.compile(r"^\[\d+\]")               # [1] 编号制
_DOI_RE = re.compile(r"doi:?\.?org|doi:\s*1\.", re.IGNORECASE)

_STYLE_LABEL = {
    "gb7714": "GB/T 7714",
    "author_year": "作者-年份制（类 APA）",
    "numbered": "编号制（顺序编码）",
    "unknown": "无法确定",
}


class FormatCheckError(Exception):
    """解档/规则加载失败（路由器转 PARSE_FAILED）。"""


# ---------------------------------------------------------------- 规则集加载

def load_ruleset(target_journal: str) -> tuple[dict[str, Any], dict[str, Any]]:
    """返回 (rules, ruleset_meta)。target_journal 传 id 或 name 均可；未知 → default。"""
    try:
        data = json.loads(_JOURNALS_PATH.read_text("utf-8"))
    except Exception as exc:  # noqa: BLE001
        raise FormatCheckError(f"规则集加载失败：{exc}") from exc

    default = data.get("format_rules_default") or {}
    jid = (target_journal or "").strip()
    entry = None
    if jid:
        for j in data.get("journals") or []:
            if j.get("id") == jid or j.get("name") == jid:
                entry = j
                break

    if entry is None:
        rules = dict(default)
        meta = {
            "journal_id": None,
            "journal_name": None,
            "ruleset_source": "default",
            "rules_source": str(default.get("rules_source", "default")),
        }
    else:
        journal_rules = entry.get("format_rules") or {}
        rules = {**default, **journal_rules}
        meta = {
            "journal_id": entry.get("id"),
            "journal_name": entry.get("name"),
            "ruleset_source": "journal",
            # 占位规则（无官方核验）→ 前端黄字提示
            "rules_source": str(journal_rules.get("rules_source", "placeholder")),
        }
    rules.setdefault("rules_source", "default")
    return rules, meta


# ---------------------------------------------------------------- PDF 解档（pymupdf）

def _norm_font(name: str) -> str:
    n = (name or "").split("+")[-1].strip()
    return n or "unknown"


def extract_profile_pdf(data: bytes) -> dict[str, Any]:
    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:  # noqa: BLE001
        raise FormatCheckError(f"PDF 解档失败：{exc}") from exc
    if getattr(doc, "needs_pass", False):
        doc.close()
        raise FormatCheckError("PDF 已加密，无法解析")

    pages_total = doc.page_count
    scan = min(pages_total, _MAX_SCAN_PAGES)
    margin_rows: list[tuple[float, float, float, float]] = []  # t r b l (mm)
    font_w: dict[str, int] = {}
    size_w: dict[float, int] = {}
    ratios: list[float] = []
    img_ids: set[int] = set()
    fig_caps = tab_caps = 0

    for pno in range(scan):
        page = doc[pno]
        rect = page.rect
        blocks = page.get_text("blocks")
        content = [b for b in blocks if b[4].strip()]
        if content:
            left = min(b[0] for b in content)
            top = min(b[1] for b in content)
            right = rect.width - max(b[2] for b in content)
            bottom = rect.height - max(b[3] for b in content)
            margin_rows.append(
                (
                    max(0.0, top) * _PT_TO_MM,
                    max(0.0, right) * _PT_TO_MM,
                    max(0.0, bottom) * _PT_TO_MM,
                    max(0.0, left) * _PT_TO_MM,
                )
            )

        for img in page.get_images(full=True):
            img_ids.add(img[0])

        text = page.get_text()
        fig_caps += len(_FIG_RE.findall(text))
        tab_caps += len(_TAB_RE.findall(text))

        d = page.get_text("dict")
        for block in d.get("blocks", []):
            if block.get("type") != 0:
                continue
            line_meta: list[tuple[float, float]] = []  # (基线 y, 字号)
            for line in block.get("lines", []):
                spans = [s for s in line.get("spans", []) if s.get("text", "").strip()]
                if not spans:
                    continue
                row_size = max(s.get("size", 0.0) for s in spans)
                if row_size <= 0:
                    continue
                for s in spans:
                    size = round(float(s.get("size", 0.0)), 1)
                    chars = len(s.get("text", ""))
                    if size > 0:
                        norm = _norm_font(s.get("font", ""))
                        font_w[norm] = font_w.get(norm, 0) + chars
                        size_w[size] = size_w.get(size, 0) + chars
                # 行距基准用 span.origin 基线 y，不用 bbox 顶缘：bbox 顶随字形高低
                # 漂移，表格/图注等同基线并列文本会产出假性小 delta，混入中位数会把
                # 行距拉成 0.1x 的不可能值。
                origin = spans[0].get("origin")
                line_meta.append((origin[1] if origin else line["bbox"][1], row_size))
            for i in range(1, len(line_meta)):
                delta = line_meta[i][0] - line_meta[i - 1][0]
                size = line_meta[i][1]
                # 行距倍数 = 相邻行基线 y 差 / 当前行字号；
                # 过滤物理不可能值：<0.5（行重叠/同基线并列表格）、>4.0（异常结构）
                ratio = delta / size
                if 0.5 <= ratio <= 4.0:
                    ratios.append(ratio)

    refs_text, refs_found = _find_refs_pdf(doc, pages_total)
    doc.close()

    if margin_rows:
        med = [statistics.median(col) for col in zip(*margin_rows)]
        margins = {"top": med[0], "right": med[1], "bottom": med[2], "left": med[3]}
    else:
        margins = None

    return {
        "source_kind": "pdf",
        "pages": pages_total,
        "margins_mm": margins,
        "font_weights": font_w,
        "size_weights": size_w,
        "line_spacing_median": statistics.median(ratios) if ratios else None,
        "figures": max(len(img_ids), fig_caps),
        "tables": tab_caps,
        "refs_text": refs_text,
        "refs_found": refs_found,
    }


def _find_refs_pdf(doc: "fitz.Document", pages_total: int) -> tuple[str | None, bool]:
    """从文末向前找参考文献标题行；找到 → 标题行至文末文本；找不到 → (None, False)。"""
    buf: list[str] = []
    start = max(0, pages_total - _MAX_REFS_SCAN)
    for pno in range(pages_total - 1, start - 1, -1):
        page_text = doc[pno].get_text()
        lines = page_text.splitlines()
        for idx in range(len(lines) - 1, -1, -1):
            if _HEADING_RE.match(lines[idx].strip()):
                tail = "\n".join(lines[idx + 1:])
                joined = "\n".join([tail] + buf).strip()
                return (joined if len(joined) >= 50 else None), bool(joined)
        buf.insert(0, page_text)
    return None, False


# ---------------------------------------------------------------- Word 解档（python-docx）

def _docx_app_pages(data: bytes) -> int | None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            xml = zf.read("docProps/app.xml").decode("utf-8", "ignore")
        m = re.search(r"<Pages>(\d+)</Pages>", xml)
        return int(m.group(1)) if m else None
    except Exception:  # noqa: BLE001
        return None


def extract_profile_docx(data: bytes) -> dict[str, Any]:
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise FormatCheckError(f"Word 文件解析失败：{exc}") from exc

    side_lists: dict[str, list[float]] = {"top": [], "right": [], "bottom": [], "left": []}
    for sec in doc.sections:
        for k, v in (
            ("top", sec.top_margin),
            ("right", sec.right_margin),
            ("bottom", sec.bottom_margin),
            ("left", sec.left_margin),
        ):
            if v is not None:
                side_lists[k].append(float(v) * _EMU_TO_MM)
    margins = None
    if all(side_lists.values()):
        margins = {k: statistics.median(v) for k, v in side_lists.items()}

    try:
        normal = doc.styles["Normal"].font
        default_name = normal.name
        default_size = float(normal.size.pt) if normal.size is not None else None
    except Exception:  # noqa: BLE001
        default_name, default_size = None, None

    font_w: dict[str, int] = {}
    size_w: dict[float, int] = {}
    ratios: list[float] = []
    fig_caps = tab_caps = 0
    paras = doc.paragraphs
    scan_count = min(len(paras), _MAX_PARAS)

    for i in range(scan_count):
        para = paras[i]
        text = para.text
        if not text.strip():
            continue
        ls = para.paragraph_format.line_spacing
        if isinstance(ls, Length):
            # 精确行距（exactly）以 EMU 存储，Length 是 int 子类——直接当倍数会得出
            # 十万级的假值；用正文默认字号换算成倍数
            if default_size:
                r = float(ls.pt) / default_size
                if 0.5 <= r <= 4.0:
                    ratios.append(r)
        elif isinstance(ls, (int, float)) and not isinstance(ls, bool) and ls > 0:
            r = float(ls)  # multiple 行距本身就是倍数
            if 0.5 <= r <= 4.0:
                ratios.append(r)
        if _FIG_RE.match(text):
            fig_caps += 1
        elif _TAB_RE.match(text):
            tab_caps += 1
        for run in para.runs:
            if not run.text:
                continue
            name = run.font.name or default_name or "unknown"
            pt = run.font.size.pt if run.font.size is not None else default_size
            chars = len(run.text)
            font_w[str(name)] = font_w.get(str(name), 0) + chars
            if pt:
                size_w[round(float(pt), 1)] = size_w.get(round(float(pt), 1), 0) + chars

    # 参考文献节：标题段落之后全部文本
    refs_text, refs_found = None, False
    for i, para in enumerate(paras):
        if _HEADING_RE.match((para.text or "").strip()):
            tail = "\n".join(p.text for p in paras[i + 1:] if (p.text or "").strip())
            if len(tail) >= 50:
                refs_text, refs_found = tail, True
            break

    return {
        "source_kind": "docx",
        "pages": _docx_app_pages(data),  # Word 未记录 → None（检查项降级为 warn）
        "margins_mm": margins,
        "font_weights": font_w,
        "size_weights": size_w,
        "line_spacing_median": statistics.median(ratios) if ratios else None,
        "figures": max(len(doc.inline_shapes), fig_caps),
        "tables": max(len(doc.tables), tab_caps),
        "refs_text": refs_text,
        "refs_found": refs_found,
    }


# ---------------------------------------------------------------- LaTeX 解档（正则+状态机，零新依赖）
#
# 决策前提：只做"解析 + 检查 + 建议代码片段"，不自动改写 LaTeX（用户可能有自定义宏，
# 自动改写会引错；LaTeX 应由用户自己编辑）。不引入 pylatexenc/texsoup。

# 会相互覆盖的字体类宏包：≥2 个同时引入 → warn（times + mathptmx 类冲突）
_TEX_FONT_PKGS = {"times", "mathptmx", "newtx", "newpx", "txfonts", "pxfonts", "mathpazo"}


def _strip_tex_comments(src: str) -> str:
    r"""去行内 % 注释（r-string，避免 \% 被解释为转义序列）。"""
    return re.sub(r"(?<!\\)%.*", "", src)


def _parse_geometry(src: str) -> dict[str, float] | None:
    """\\usepackage[...]{geometry} 或 \\geometry{...} → {边: mm}（可缺边，检查端容缺）。"""
    m = re.search(r"\\usepackage\s*\[([^\]]*)\]\s*\{\s*geometry\s*\}", src)
    if not m:
        m = re.search(r"\\geometry\s*\{([^}]*)\}", src)
        if not m:
            return None
    opts = m.group(1)

    def to_mm(val: str) -> float | None:
        mm = re.match(r"\s*([\d.]+)\s*(mm|cm|in|pt)?\s*$", val)
        if not mm:
            return None
        x = float(mm.group(1))
        unit = mm.group(2) or "pt"  # geometry 裸数字按 TeX 惯例为 pt
        return {"mm": x, "cm": x * 10, "in": x * 25.4, "pt": x * 25.4 / 72.0}[unit]

    kv = {}
    for pair in opts.split(","):
        if "=" in pair:
            k, v = pair.split("=", 1)
            kv[k.strip().lower()] = v.strip()
    out: dict[str, float] = {}
    if "margin" in kv:
        v = to_mm(kv["margin"])
        if v is not None:
            return {"top": v, "right": v, "bottom": v, "left": v}
    for k in ("top", "right", "bottom", "left"):
        if k in kv:
            v = to_mm(kv[k])
            if v is not None:
                out[k] = v
    if "inner" in kv:
        v = to_mm(kv["inner"])
        if v is not None:
            out.setdefault("left", v)
    if "outer" in kv:
        v = to_mm(kv["outer"])
        if v is not None:
            out.setdefault("right", v)
    return out or None


def extract_tex_profile(data: bytes) -> dict[str, Any]:
    """正则 + 轻量状态机解析 .tex 源文件 → 与 DocProfile 同构（+tex 扩展字段）。"""
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode("latin-1", errors="replace")
    src = _strip_tex_comments(text)

    # \documentclass[10pt,a4paper]{article}
    m_cls = re.search(r"\\documentclass\s*(?:\[[^\]]*\])?\s*\{\s*([\w*]+)\s*\}", src)
    cls = m_cls.group(1) if m_cls else None
    m_opt = re.search(r"\\documentclass\s*\[([^\]]*)\]", src)
    cls_opts = m_opt.group(1) if m_opt else ""
    m_pt = re.search(r"(\d+)\s*pt", cls_opts)
    # 标准类未写选项时 LaTeX 默认 10pt；无 documentclass → None（检查项降级 warn）
    fontsize = int(m_pt.group(1)) if m_pt else (10 if cls else None)
    paper = "a4paper" if "a4paper" in cls_opts else ("letterpaper" if "letterpaper" in cls_opts else None)

    # \usepackage 列表（支持 {a,b} 逗号多包）
    packages: list[str] = []
    for m in re.finditer(r"\\usepackage\s*(?:\[[^\]]*\])?\s*\{\s*([^}]+?)\s*\}", src):
        packages += [p.strip().lower() for p in m.group(1).split(",") if p.strip()]

    # 行距：\linespread{...} 或 \setstretch{...}（值直接按"行距倍数"口径评估——
    # 按步骤4 规格：低于 min_line_spacing 即 fail；LaTeX 实际倍数还含 baselinestretch
    # 的字号比，本引擎取显式设置值做保守判定）
    line_spread: float | None = None
    m = re.search(r"\\linespread\s*\{\s*([\d.]+)\s*\}", src)
    if not m:
        m = re.search(r"\\setstretch\s*\{\s*([\d.]+)\s*\}", src)
    if m:
        line_spread = float(m.group(1))

    # 页边距（geometry）
    margins = _parse_geometry(src)

    # 计数
    figures = len(re.findall(r"\\begin\s*\{\s*figure\*?\s*\}", src))
    tables = len(re.findall(r"\\begin\s*\{\s*table\*?\s*\}", src))
    cites = len(re.findall(r"\\cite\w*\s*\{", src))

    # 参考文献：\bibliography{...}（BibTeX）/ \printbibliography（biblatex）/ thebibliography
    has_bibtex = bool(re.search(r"\\bibliography\s*\{", src))
    has_printbib = bool(re.search(r"\\printbibliography", src))
    has_thebib = bool(re.search(r"\\begin\s*\{\s*thebibliography\s*\}", src))
    has_bibstyle = bool(re.search(r"\\bibliographystyle\s*\{", src))  # BibTeX 必需；biblatex 不需要
    has_bib = has_bibtex or has_printbib or has_thebib
    refs_text: str | None = None
    m_bib = re.search(r"\\begin\s*\{\s*thebibliography\s*\}[\s\S]*?\\end\s*\{\s*thebibliography\s*\}", src)
    if m_bib:
        refs_text = m_bib.group(0)

    # 宏表（\newcommand{\xxx} / \renewcommand）——用于解析准确性（识别自定义宏）
    macros = re.findall(r"\\(?:new|re)command\s*\{\s*\\(\w+)\s*\}", src)

    # 匿名评审命令探测（规则集 anonymous:true 时检查）
    m_anon = re.search(r"\\(?:anonymi[sz]e|anon\b|hideauthor|ifdoubleblind)", src, re.IGNORECASE)

    return {
        "source_kind": "tex",
        "pages": None,  # LaTeX 源文件页数需编译后才有 → 检查项降级 warn（宽容化）
        "margins_mm": margins,
        "font_weights": {},
        "size_weights": {float(fontsize): 100} if fontsize else {},
        "line_spacing_median": line_spread,
        "figures": figures,
        "tables": tables,
        "refs_text": refs_text,
        "refs_found": has_bib,
        "tex": {
            "class": cls,
            "fontsize_pt": fontsize,
            "paper": paper,
            "packages": packages,
            "cites": cites,
            "macros": macros,
            "anon_cmd": m_anon.group(0) if m_anon else None,
            "bib_bibtex": has_bibtex,
            "bib_style": has_bibstyle,
        },
    }


# ---------------------------------------------------------------- 参考文献样式检测

def detect_refs_style(text: str | None) -> tuple[str, str]:
    """宽容化检测：返回 (detected, sample_line)。含糊 → unknown（后续只 warn）。"""
    if not text:
        return "unknown", ""
    lines = [l.strip() for l in text.splitlines() if l.strip()][:40]
    if not lines:
        return "unknown", ""
    n = len(lines)
    gb = sum(1 for l in lines if _GB_TAG_RE.search(l))
    year = sum(1 for l in lines if _PAREN_YEAR_RE.search(l))
    num = sum(1 for l in lines if _BRACKET_NUM_RE.match(l))
    doi = sum(1 for l in lines if _DOI_RE.search(l))
    sample = lines[0][:80]

    if gb / n >= 0.4:
        return "gb7714", sample
    if num / n >= 0.4 and year / n < 0.3:
        return "numbered", sample
    if year / n >= 0.3 or doi / n >= 0.3:
        return "author_year", sample
    return "unknown", sample


# ---------------------------------------------------------------- 检查项构建

def _chk(key: str, name: str, status: str, detail: str, suggestion: str = "") -> dict[str, str]:
    return {"key": key, "name": name, "status": status, "detail": detail, "suggestion": suggestion}


def _check_pages(profile: dict[str, Any], rules: dict[str, Any]) -> dict[str, str]:
    if rules.get("max_pages") is None:
        return _chk("pages", "页数", "warn", "目标期刊未给出统一页数上限，请按稿件类型核对")
    limit = int(rules["max_pages"])
    pages = profile.get("pages")
    kind = profile.get("source_kind") or "doc"
    if pages is None:
        # 宽容化：Word/LaTeX 页数缺失绝不 fail
        if kind == "tex":
            return _chk(
                "pages", "页数", "warn",
                f"LaTeX 源无法直接统计页数（规则上限 {limit} 页），需编译后核对",
                "编译后确认（如 xelatex main.tex）总页数 ≤ 上限",
            )
        return _chk(
            "pages", "页数", "warn",
            "Word 未记录页数，请自行核对",
            f"在 Word 中按 Ctrl+End 查看总页数，或导出 PDF 后重查（规则上限 {limit} 页）",
        )
    if pages <= limit:
        return _chk("pages", "页数", "pass", f"共 {pages} 页，规则上限 {limit} 页")
    return _chk(
        "pages", "页数", "fail", f"共 {pages} 页，超出规则上限 {limit} 页",
        f"压缩至 ≤{limit} 页，或将方法/图表移入补充材料",
    )


def _check_fonts(profile: dict[str, Any], rules: dict[str, Any]) -> dict[str, str]:
    min_font = float(rules.get("min_font_pt", 9))
    # LaTeX：字号来自 \documentclass 选项（10/11/12pt，粒度粗）
    if profile.get("source_kind") == "tex":
        tex = profile.get("tex") or {}
        cls = tex.get("class") or "documentclass"
        fs = tex.get("fontsize_pt")
        pkgs = tex.get("packages") or []
        font_pkgs = sorted(p for p in pkgs if p in _TEX_FONT_PKGS or p in ("helvet", "courier", "charter", "fourier"))
        detail = (
            f"\\documentclass 选项字号 {fs}pt（类 {cls}）；"
            f"字体宏包：{'、'.join(font_pkgs) if font_pkgs else '未引入'}"
            + (f"；纸张 {tex['paper']}" if tex.get("paper") else "")
        )
        if fs is None:
            return _chk("fonts", "字体字号", "warn", detail + "\n建议：\n  \\documentclass[11pt]{article}", "补充 \\documentclass 声明")
        if fs < min_font:
            # LaTeX 粒度粗 → 按步骤4 规格降为 warn（非 fail）
            return _chk(
                "fonts", "字体字号", "warn", detail + f"\n建议：\n  \\documentclass[11pt]{{{cls}}}",
                f"正文字号低于下限 {min_font:.0f}pt，请改用 11pt/12pt 选项",
            )
        return _chk("fonts", "字体字号", "pass", detail)

    size_w: dict[float, int] = profile.get("size_weights") or {}
    font_w: dict[str, int] = profile.get("font_weights") or {}
    if not size_w:
        return _chk("fonts", "字体字号", "warn", "未能提取字体信息", "确认文档含可选中文本（扫描件无法检查字体）")
    min_font = float(rules.get("min_font_pt", 9))
    max_fonts = int(rules.get("max_fonts", 6))
    total = sum(size_w.values())
    body_pt = max(size_w, key=size_w.get)
    body_ratio = int(size_w[body_pt] / total * 100)
    distinct = len(font_w)
    top_fonts = sorted(font_w.items(), key=lambda kv: -kv[1])[:3]
    top_text = "、".join(f"{k}({v}字)" for k, v in top_fonts)
    detail = f"主要字号 {body_pt}pt（约 {body_ratio}%），共 {distinct} 种字体；高频字体：{top_text}"
    if body_pt < min_font:
        return _chk(
            "fonts", "字体字号", "fail", detail,
            f"将正文字号提高至 ≥{min_font}pt（当前主字号 {body_pt}pt）",
        )
    if distinct > max_fonts:
        return _chk(
            "fonts", "字体字号", "warn", detail,
            f"字体种类偏多（{distinct} > {max_fonts}），建议全文统一为 1–2 种正文字体",
        )
    return _chk("fonts", "字体字号", "pass", detail)


def _check_margins(profile: dict[str, Any], rules: dict[str, Any]) -> dict[str, str]:
    limit = float(rules.get("min_margin_mm", 15))
    margins = profile.get("margins_mm")
    # LaTeX：页边距来自 geometry；低于下限按步骤4 规格给 warn（非 fail），附建议片段
    if profile.get("source_kind") == "tex":
        if not margins:
            return _chk(
                "margins", "页边距", "warn",
                "未通过 geometry 设置页边距（按类默认评估）"
                + f"\n建议：\n  \\usepackage[margin=25mm]{{geometry}}",
                "在导言区显式声明页边距",
            )
        lo = min(margins.values())
        detail = f"geometry 页边距最小 {lo:.0f}mm（" + "、".join(f"{k}{v:.0f}" for k, v in margins.items()) + f"），规则下限 {limit:.0f}mm"
        if lo < limit:
            return _chk(
                "margins", "页边距", "warn",
                detail + f"\n建议：\n  \\usepackage[margin={limit:.0f}mm]{{geometry}}",
                f"调大页边距至 ≥{limit:.0f}mm",
            )
        return _chk("margins", "页边距", "pass", detail)
    if not margins:
        return _chk("margins", "页边距", "warn", "未能提取页边距信息", "检查文档是否含可定位文本")
    vals = [margins.get(k) for k in ("top", "right", "bottom", "left")]
    vals = [v for v in vals if v is not None]
    if not vals:
        return _chk("margins", "页边距", "warn", "未能提取页边距信息", "检查文档是否含可定位文本")
    t, r = margins.get("top"), margins.get("right")
    b, l = margins.get("bottom"), margins.get("left")
    lo = min(vals)
    sides = "、".join(
        f"{name}{v:.0f}" for name, v in (("上", t), ("右", r), ("下", b), ("左", l)) if v is not None
    )
    detail = f"四边最小 {lo:.0f}mm（{sides}），规则下限 {limit:.0f}mm"
    if lo >= limit:
        return _chk("margins", "页边距", "pass", detail)
    if lo >= limit * 0.8:
        return _chk("margins", "页边距", "warn", detail, f"将页边距调大至 ≥{limit:.0f}mm")
    return _chk("margins", "页边距", "fail", detail, f"页边距过小，请调至 ≥{limit:.0f}mm（版式设置）")


def _check_spacing(profile: dict[str, Any], rules: dict[str, Any]) -> dict[str, str]:
    lo = float(rules.get("min_line_spacing", 1.15))
    ratio = profile.get("line_spacing_median")
    # LaTeX：\linespread/\setstretch 显式设置值；低于下限按步骤4 规格直接 fail
    if profile.get("source_kind") == "tex":
        if ratio is None:
            return _chk(
                "spacing", "行距", "pass",
                "未显式设置 \\linespread/\\setstretch，LaTeX 类默认约 1.2 倍（如 10pt→12pt baselineskip）",
            )
        detail = f"\\linespread/\\setstretch = {ratio:.2f}，规则下限 {lo:.2f}"
        if ratio < lo:
            return _chk(
                "spacing", "行距", "fail",
                detail + f"\n建议：\n  \\linespread{{{lo:.2f}}}   % 或 \\setstretch{{{lo:.2f}}}（导言区）",
                "在导言区调大行距设置",
            )
        return _chk("spacing", "行距", "pass", detail)
    if ratio is None:
        return _chk("spacing", "行距", "warn", "未能提取行距信息", "确认段落使用固定/多倍行距设置")
    lo = float(rules.get("min_line_spacing", 1.15))
    detail = f"正文中位行距约 {ratio:.2f} 倍，规则下限 {lo:.2f} 倍"
    if ratio >= lo:
        return _chk("spacing", "行距", "pass", detail)
    if ratio >= lo * 0.9:
        return _chk("spacing", "行距", "warn", detail, f"段落行距设为 ≥{lo:.2f} 倍（如 1.5 倍行距）")
    return _chk("spacing", "行距", "fail", detail, f"行距过密，请设为 ≥{lo:.2f} 倍行距")


def _check_figures(profile: dict[str, Any], rules: dict[str, Any]) -> dict[str, str]:
    figures = int(profile.get("figures") or 0)
    tables = int(profile.get("tables") or 0)
    min_f = int(rules.get("min_figures", 1))
    detail = f"检测到图 {figures} 张、表 {tables} 个（图注/表注正则 + 内嵌对象启发式）"
    if figures >= min_f:
        return _chk("figures", "图表数量", "pass", detail)
    # 启发式检测不可靠 → 低于阈值只 warn 不 fail
    return _chk(
        "figures", "图表数量", "warn", detail,
        f"常规研究论文建议 ≥{min_f} 图；若为综述/短文可忽略（本项为启发式检测）",
    )


# 仅"明确不符"才 fail 的 (expected, detected) 组合（宽容化：其余含糊组合一律 warn）
_REF_FAIL_PAIRS = {
    ("apa", "gb7714"),        # 找到 [J] 标引但规则要求 APA
    ("gb7714", "author_year"),  # 无 [J] 标引的作者-年份制，但规则要求 GB/T 7714
    ("apa", "numbered"),      # 编号制引文但规则要求 APA
}


def _check_refs(profile: dict[str, Any], rules: dict[str, Any]) -> dict[str, str]:
    expected = str(rules.get("refs_style", "apa")).lower()
    exp_label = _STYLE_LABEL.get(expected, expected)
    if not profile.get("refs_found"):
        return _chk(
            "refs", "参考文献格式", "warn",
            "未找到参考文献节（References/参考文献 / \\bibliography / thebibliography）",
            "确认存在参考文献声明后重查",
        )
    if profile.get("source_kind") == "tex":
        tex = profile.get("tex") or {}
        # BibTeX：\bibliography 存在但缺 \bibliographystyle → warn（样式引擎未知）
        if tex.get("bib_bibtex") and not tex.get("bib_style"):
            return _chk(
                "refs", "参考文献格式", "warn",
                f"规则要求 {exp_label}；检测到 \\bibliography 但缺少 \\bibliographystyle（BibTeX 无法确定引用样式）"
                + "\n建议：\n  \\bibliographystyle{plain}   % 按期刊/模板指定实际样式",
                "在 \\bibliography 前补 \\bibliographystyle",
            )
        if not profile.get("refs_text"):
            # 条目在 .bbl/编译产物，源文件无文本 → 宽容化只 warn
            return _chk(
                "refs", "参考文献格式", "warn",
                f"规则要求 {exp_label}；检测到参考文献声明（条目由编译生成，源文件无条目文本）",
                "编译后核对参考文献条目样式",
            )
    detected, sample = detect_refs_style(profile.get("refs_text"))
    det_label = _STYLE_LABEL[detected]
    detail = f"规则要求 {exp_label}；检测到 {det_label}；样例：{sample}"
    # 宽容化：仅"明确不符"（_REF_FAIL_PAIRS）才 fail；含糊/不确定一律 warn
    if (expected, detected) in _REF_FAIL_PAIRS:
        return _chk("refs", "参考文献格式", "fail", detail, f"按 {exp_label} 调整全部文献条目（可用文献管理工具一键转换）")
    if detected == "unknown":
        return _chk("refs", "参考文献格式", "warn", detail, f"请人工核对是否符合 {exp_label}")
    if expected == "gb7714" and detected == "numbered":
        return _chk("refs", "参考文献格式", "warn", detail, "GB/T 7714 顺序编码制也用编号，需人工确认条目含 [J] 等类型标引")
    if expected == "author-year-journal" and detected in ("gb7714", "numbered"):
        return _chk("refs", "参考文献格式", "warn", detail, f"请对照期刊投稿指南核对是否符合 {exp_label}")
    return _chk("refs", "参考文献格式", "pass", detail)


def _check_font_packages(profile: dict[str, Any], rules: dict[str, Any]) -> dict[str, str]:
    """LaTeX 特有：多个相互覆盖的字体宏包（times + mathptmx 类）→ warn。"""
    pkgs = set((profile.get("tex") or {}).get("packages") or [])
    clash = sorted(pkgs & _TEX_FONT_PKGS)
    if len(clash) >= 2:
        detail = (
            f"检测到多个相互覆盖的字体宏包：{'、'.join(clash)}（后加载者覆盖前者，易产生冲突）"
            + "\n建议：\n  %只保留一套字体方案，例如：\n  \\usepackage{newtx}   % newtxtext + newtxmath 一体"
        )
        return _chk("fontpkgs", "字体包兼容", "warn", detail, "移除重复/冲突的字体宏包，只保留一套")
    return _chk(
        "fontpkgs", "字体包兼容", "pass",
        "字体宏包：" + ("、".join(clash) if clash else "未引入冲突类字体包"),
    )


def _check_anonymous(profile: dict[str, Any], rules: dict[str, Any]) -> dict[str, str]:
    """LaTeX 特有且规则集 anonymous:true 时才出：源文件需含匿名处理命令。"""
    tex = profile.get("tex") or {}
    cmd = tex.get("anon_cmd")
    if cmd:
        return _chk("anonymous", "匿名评审", "pass", f"检测到匿名处理命令 {cmd}")
    return _chk(
        "anonymous", "匿名评审", "warn",
        "规则要求匿名评审，源文件未检测到匿名处理命令"
        + "\n建议：\n  %隐去作者/致谢/自引，或定义匿名宏：\n  \\newcommand{\\anon}[1]{[ANON]}",
        "按期刊要求对作者信息做匿名处理",
    )


def build_checks(profile: dict[str, Any], rules: dict[str, Any]) -> list[dict[str, str]]:
    checks = []
    for key, label, fields, check in [
        ("pages", "页数", ["max_pages"], _check_pages),
        ("fonts", "字体字号", ["min_font_pt", "max_fonts"], _check_fonts),
        ("margins", "页边距", ["min_margin_mm"], _check_margins),
        ("spacing", "行距", ["min_line_spacing"], _check_spacing),
        ("figures", "图表", ["min_figures"], _check_figures),
    ]:
        if any(field in rules and rules[field] is None for field in fields):
            checks.append(_chk(key, label, "warn", "期刊未提供完整量化要求，请按稿件类型核对官方指南"))
        else:
            checks.append(check(profile, rules))
    if profile.get("source_kind") == "tex":
        checks.append(_check_font_packages(profile, rules))
    checks.append(_check_refs(profile, rules))
    if profile.get("source_kind") == "tex" and rules.get("anonymous"):
        checks.append(_check_anonymous(profile, rules))
    return checks


# ---------------------------------------------------------------- 入口

def check_document(data: bytes, kind: str, target_journal: str) -> dict[str, Any]:
    """kind ∈ pdf|docx|tex。返回 {checks, summary, meta}（filename 由路由器补入 meta）。"""
    if kind == "pdf":
        profile = extract_profile_pdf(data)
    elif kind == "docx":
        profile = extract_profile_docx(data)
    elif kind == "tex":
        profile = extract_tex_profile(data)
    else:
        raise FormatCheckError(f"不支持的文件类型：{kind}")

    rules, ruleset_meta = load_ruleset(target_journal)
    checks = build_checks(profile, rules)
    summary = {"pass": 0, "warn": 0, "fail": 0}
    for c in checks:
        summary[c["status"]] = summary.get(c["status"], 0) + 1

    detected, _ = detect_refs_style(profile.get("refs_text")) if profile.get("refs_found") else ("unknown", "")
    meta = {
        **ruleset_meta,
        "file_type": kind,
        "pages": profile.get("pages"),
        "refs_style_expected": str(rules.get("refs_style", "apa")).lower(),
        "refs_style_detected": detected,
    }
    if kind == "tex":
        tex = profile.get("tex") or {}
        meta["tex"] = {
            "class": tex.get("class"),
            "fontsize_pt": tex.get("fontsize_pt"),
            "paper": tex.get("paper"),
            "packages": tex.get("packages"),
            "cites": tex.get("cites"),
            "macros": (tex.get("macros") or [])[:10],
        }
    return {"checks": checks, "summary": summary, "meta": meta}
