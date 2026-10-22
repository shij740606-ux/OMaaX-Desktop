"""app/services/journal_seed.py —— journal_directory 核心导入逻辑（CLI 与启动 seed 共享）。

数据流（schema 仍由本模块独家拥有，app 路由层只读不建表）：
  openpyxl 非 read_only 加载（读 cell.hyperlink.target 来源链接）
  → 「国际期刊」「中文期刊」两 sheet 按【列名】映射
  → 行级清洗（空→None 绝不填 0；分区数字→"N区"；JCR学科→JSON 数组）
  → 按 id 幂等 upsert（ON CONFLICT DO UPDATE，重跑 count 不变）

两个入口：
  - scripts/import_journals.py（CLI）：extract_rows_from_xlsx + upsert_to_db
    + 对账报告（exit code 语义见脚本 docstring）
  - ensure_journal_directory_seeded()：服务启动自动 seed（main.py lifespan），
    仅当表不存在或为空才导入；任何失败收敛为 dict 返回值，绝不抛异常。

v4 源（2026-10-10）：JIF=2025 统计年/2026 JCR 发布；中文 IF=2024 统计年/
年报 2025 版。源 xlsx 路径走 settings.journal_directory_seed_abs
（REVIEW_JOURNAL_DIRECTORY_SEED_PATH 覆盖）。

id 策略：有 ISSN → "{lang}:{ISSN}"; 无 ISSN → "{lang}:{slug(name)}"；
同 lang 内冲突 → 抛 JournalSeedError（CLI 中止导入 / startup 返回 error）。
"""
from __future__ import annotations

import json
import re
import sqlite3
import unicodedata
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl

from ..config import get_settings

SHEET_INTL = "国际期刊"
SHEET_ZH = "中文期刊"
EXPECT_INTL = 779  # 口径 sheet 声明（v4 不变）
EXPECT_ZH = 277

# 官网列（v4 实际列名「官网网址」在首位；其余为历史/备用别名）
OFFICIAL_URL_HEADERS = ("官网网址", "官网", "官方网址", "期刊官网")
# JCR 学科列别名：v4 改名「JCR学科（2024源表）」，v2 为「JCR学科（源表）」
JCR_DISCIPLINE_HEADERS = ("JCR学科（2024源表）", "JCR学科（源表）")

# 既有表缺列时的幂等迁移（缺哪加哪，重复跑不报错）
MIGRATE_COLS = (
    ("official_url", "TEXT"),
    ("official_check_status", "TEXT"),
    ("jcr_release_year", "INTEGER"),
)

SCHEMA_SQL = """
CREATE TABLE IF NOT EXISTS journal_directory (
    id                          TEXT PRIMARY KEY,
    lang                        TEXT NOT NULL,
    name                        TEXT NOT NULL,
    research_direction          TEXT,
    cas_partition               TEXT,
    cas_partition_note          TEXT,
    cas_partition_year          INTEGER,
    jif                        REAL,
    jif_year                    INTEGER,
    jif_note                    TEXT,
    jif_check_status            TEXT,
    composite_impact_factor     REAL,
    comprehensive_impact_factor REAL,
    metric_year                 INTEGER,
    metric_report_year          INTEGER,
    metric_caliber              TEXT,
    annual_report_page          INTEGER,
    top_flag                    TEXT,
    publisher_series            TEXT,
    journal_mark                TEXT,
    article_type                TEXT,
    issn                        TEXT,
    eissn                       TEXT,
    jcr_disciplines             TEXT,
    supplement_note             TEXT,
    evidence_note               TEXT,
    official_url                TEXT,
    official_check_status       TEXT,
    jcr_release_year            INTEGER,
    partition_source            TEXT,
    jif_source                  TEXT,
    jif_source_url              TEXT,
    partition_source_url        TEXT,
    data_source_label           TEXT,
    data_source_url             TEXT,
    imported_at                 TEXT
);
CREATE INDEX IF NOT EXISTS idx_dir_lang_name    ON journal_directory(lang, name);
CREATE INDEX IF NOT EXISTS idx_dir_lang_jif     ON journal_directory(lang, jif);
CREATE INDEX IF NOT EXISTS idx_dir_lang_cif     ON journal_directory(lang, composite_impact_factor);
CREATE INDEX IF NOT EXISTS idx_dir_lang_part    ON journal_directory(lang, cas_partition);
CREATE INDEX IF NOT EXISTS idx_dir_lang_domain  ON journal_directory(lang, research_direction);
"""

# upsert 列序（除 id 外全部 ON CONFLICT DO UPDATE）
COLS = (
    "lang", "name", "research_direction", "cas_partition", "cas_partition_note",
    "cas_partition_year", "jif", "jif_year", "jif_note", "jif_check_status",
    "composite_impact_factor", "comprehensive_impact_factor", "metric_year",
    "metric_report_year", "metric_caliber", "annual_report_page", "top_flag",
    "publisher_series", "journal_mark", "article_type", "issn", "eissn",
    "jcr_disciplines", "supplement_note", "evidence_note", "official_url",
    "official_check_status", "jcr_release_year",
    "partition_source", "jif_source", "jif_source_url", "partition_source_url",
    "data_source_label", "data_source_url", "imported_at",
)
UPSERT_SQL = (
    "INSERT INTO journal_directory (id, " + ", ".join(COLS) + ") VALUES ("
    + ", ".join("?" for _ in range(len(COLS) + 1)) + ") "
    + "ON CONFLICT(id) DO UPDATE SET "
    + ", ".join(f"{c}=excluded.{c}" for c in COLS)
)


class JournalSeedError(Exception):
    """xlsx 缺失/结构错误/口径不符/id 冲突（不写库）。

    exit_code 对齐 CLI 语义：1=id 冲突中止；2=行数与口径不符；3=文件/结构错误。
    """

    def __init__(
        self,
        message: str,
        *,
        exit_code: int = 3,
        stage: str = "读取失败",
        conflicts: list[str] | None = None,
    ) -> None:
        super().__init__(message)
        self.exit_code = exit_code
        self.stage = stage
        self.conflicts = conflicts or []


def new_stats() -> dict[str, Any]:
    """对账统计骨架（CLI 报告与解析共享）。"""
    return {
        "read_intl": 0, "read_zh": 0, "success": 0, "skipped": 0, "failed": 0,
        "jif_missing": [], "cif_missing": [], "part_blank": [], "part_notfound": [],
        "part_dist": Counter(), "hl_jif": 0, "hl_part": 0, "hl_zh": 0,
        "id_issn": 0, "id_slug": 0, "official_url_detected": False,
        "migrated_cols": [], "migrated_skipped": [],
        "official_total": 0, "official_filled": 0, "official_check_filled": 0,
        "jcr_filled": 0, "official_missing": [],
    }


def migrate(conn: sqlite3.Connection) -> tuple[list[str], list[str]]:
    """幂等加列：PRAGMA 探测现有列，缺哪个 ALTER ADD 哪个。"""
    existing = {r[1] for r in conn.execute("PRAGMA table_info(journal_directory)")}
    added: list[str] = []
    skipped: list[str] = []
    for col, decl in MIGRATE_COLS:
        if col in existing:
            skipped.append(col)
        else:
            conn.execute(f"ALTER TABLE journal_directory ADD COLUMN {col} {decl}")
            added.append(col)
    conn.commit()
    return added, skipped


# ---- 清洗原语 ----

def text(v: Any) -> str | None:
    """单元格 → strip 后的文本；空/纯空白 → None（绝不填 0/空串）。"""
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def to_float(v: Any) -> float | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def to_int(v: Any) -> int | None:
    f = to_float(v)
    if f is None:
        return None
    return int(f)


def slugify(name: str) -> str:
    """小写 + 非词字符折叠为连字符；\\w 含 CJK（中文期刊无 ISSN 时的 id 兜底）。"""
    s = unicodedata.normalize("NFC", name.strip().casefold())
    s = re.sub(r"[^\w]+", "-", s, flags=re.UNICODE)
    s = re.sub(r"-{2,}", "-", s)
    return s.strip("-")


def make_id(lang: str, name: str, issn: str | None) -> str:
    if issn:
        return f"{lang}:{issn}"
    return f"{lang}:{slugify(name)}"


def parse_partition(raw: Any) -> tuple[str | None, str | None]:
    """分区 → (cas_partition, note)。数字 1-4 → 'N区'；'未检得' → note；空白 → 双 None。"""
    if raw is None:
        return None, None
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        n = int(raw)
        if 1 <= n <= 4:
            return f"{n}区", None
        return None, None
    s = str(raw).strip()
    if not s:
        return None, None
    if "未检得" in s:
        return None, "未检得"
    m = re.fullmatch(r"([1-4])\s*区?", s)
    if m:
        return f"{m.group(1)}区", None
    return s, None  # 未知文本：原样保留（不丢信息）


def split_disciplines(raw: Any) -> str | None:
    """'A(SCIE);B(SCIE)' → JSON 数组字符串。"""
    s = text(raw)
    if not s:
        return None
    parts = [p.strip() for p in s.split(";") if p.strip()]
    return json.dumps(parts, ensure_ascii=False) if parts else None


def hyperlink_target(cell) -> str | None:
    """超链接单元格 → target URL；无链接 → None（需要非 read_only 模式）。"""
    try:
        target = cell.hyperlink.target if cell.hyperlink is not None else None
    except Exception:  # noqa: BLE001
        return None
    return text(target)


# ---- sheet 解析 ----

def header_index(ws) -> dict[str, int]:
    headers = [text(c.value) for c in ws[1]]
    return {h: i for i, h in enumerate(headers) if h is not None}


def read_intl(ws, imported_at: str, stats: dict) -> list[tuple]:
    """国际期刊 sheet → 行元组（顺序同 COLS 前置 id）。"""
    ix = header_index(ws)
    for req in ("期刊名称", "研究方向", "中科院分区", "JIF影响因子"):
        if req not in ix:
            raise JournalSeedError(
                f"{SHEET_INTL} 缺列：{req}（实际列：{list(ix)}）",
                exit_code=3, stage="结构错误",
            )
    url_col = next((ix[h] for h in OFFICIAL_URL_HEADERS if h in ix), None)
    jcr_hdr = next((h for h in JCR_DISCIPLINE_HEADERS if h in ix), None)
    stats["official_url_detected"] = url_col is not None
    if url_col is not None:
        stats["official_url_headers"] = next(h for h in OFFICIAL_URL_HEADERS if h in ix)

    rows: list[tuple] = []
    for row in ws.iter_rows(min_row=2):
        if not any(c.value is not None for c in row):
            continue  # 尾部幻影空行：不计入 skipped
        name = text(row[ix["期刊名称"]].value) if ix["期刊名称"] < len(row) else None
        if not name:
            stats["skipped"] += 1
            continue
        issn = text(row[ix["ISSN"]].value) if "ISSN" in ix else None
        jif = to_float(row[ix["JIF影响因子"]].value)
        check_status = text(row[ix["JIF核对状态"]].value) if "JIF核对状态" in ix else None
        official_check = text(row[ix["官网核对状态"]].value) if "官网核对状态" in ix else None
        jcr_year = to_int(row[ix["JCR发布年"]].value) if "JCR发布年" in ix else None
        part, part_note = parse_partition(row[ix["中科院分区"]].value)

        # JIF 空 → note = 核对状态原文（"未检得"/"源表标N/A" 等说明）
        jif_note = check_status if jif is None else None

        jif_url = hyperlink_target(row[ix["JIF数据来源"]]) if "JIF数据来源" in ix else None
        part_url = hyperlink_target(row[ix["分区数据来源"]]) if "分区数据来源" in ix else None
        official = None
        if url_col is not None and url_col < len(row):
            official = text(row[url_col].value) or hyperlink_target(row[url_col])

        rows.append((
            make_id("en", name, issn), "en", name,
            text(row[ix["研究方向"]].value),
            part, part_note,
            to_int(row[ix["分区版本年"]].value) if "分区版本年" in ix else None,
            jif,
            to_int(row[ix["JIF统计年"]].value) if "JIF统计年" in ix else None,
            jif_note, check_status,
            None, None, None, None, None, None,
            text(row[ix["Top（源表）"]].value) if "Top（源表）" in ix else None,
            text(row[ix["出版系列"]].value) if "出版系列" in ix else None,
            text(row[ix["期刊标记"]].value) if "期刊标记" in ix else None,
            text(row[ix["文章类型"]].value) if "文章类型" in ix else None,
            issn,
            text(row[ix["eISSN"]].value) if "eISSN" in ix else None,
            split_disciplines(row[ix[jcr_hdr]].value) if jcr_hdr else None,
            text(row[ix["补充说明"]].value) if "补充说明" in ix else None,
            text(row[ix["补充证据"]].value) if "补充证据" in ix else None,
            official,
            official_check,
            jcr_year,
            "official", "official",
            jif_url, part_url,
            text(row[ix["JIF数据来源"]].value) if "JIF数据来源" in ix else None,
            None,  # data_source_url（国际表来源即 jif/partition 两链接）
            imported_at,
        ))
        # 统计
        stats["official_total"] += 1
        stats["official_filled"] += 1 if official else 0
        stats["official_check_filled"] += 1 if official_check else 0
        stats["jcr_filled"] += 1 if jcr_year else 0
        if not official:
            stats["official_missing"].append(("en", name))
        if jif is None:
            stats["jif_missing"].append((name, check_status or ""))
        if part is None and part_note == "未检得":
            stats["part_notfound"].append(name)
        elif part is None:
            stats["part_blank"].append(name)
        stats["hl_jif"] += 1 if jif_url else 0
        stats["hl_part"] += 1 if part_url else 0
        stats["id_issn" if issn else "id_slug"] += 1
    return rows


def read_zh(ws, imported_at: str, stats: dict) -> list[tuple]:
    """中文期刊 sheet → 行元组。"""
    ix = header_index(ws)
    for req in ("中文期刊名称", "复合影响因子", "综合影响因子"):
        if req not in ix:
            raise JournalSeedError(
                f"{SHEET_ZH} 缺列：{req}（实际列：{list(ix)}）",
                exit_code=3, stage="结构错误",
            )
    url_col = next((ix[h] for h in OFFICIAL_URL_HEADERS if h in ix), None)
    stats["official_url_detected_zh"] = url_col is not None

    rows: list[tuple] = []
    for row in ws.iter_rows(min_row=2):
        if not any(c.value is not None for c in row):
            continue  # 尾部幻影空行：不计入 skipped
        name = text(row[ix["中文期刊名称"]].value)
        if not name:
            stats["skipped"] += 1
            continue
        composite = to_float(row[ix["复合影响因子"]].value)
        comprehensive = to_float(row[ix["综合影响因子"]].value)
        check_status = text(row[ix["核对状态"]].value) if "核对状态" in ix else None

        src_label = text(row[ix["数据来源"]].value) if "数据来源" in ix else None
        src_url = hyperlink_target(row[ix["数据来源"]]) if "数据来源" in ix else None
        official = None
        if url_col is not None and url_col < len(row):
            official = text(row[url_col].value) or hyperlink_target(row[url_col])
        official_check = text(row[ix["官网核对状态"]].value) if "官网核对状态" in ix else None

        rows.append((
            make_id("zh", name, None), "zh", name,
            text(row[ix["研究方向"]].value),
            # 中文分区原样（常量"不按此口径列示"），不分区 → partition_source=None
            text(row[ix["中科院分区"]].value) if "中科院分区" in ix else None,
            None,  # cas_partition_note：中文表无"未检得"语义（口径=不按此口径列示）
            None,  # cas_partition_year
            None, None, None, check_status,
            composite, comprehensive,
            to_int(row[ix["指标统计年"]].value) if "指标统计年" in ix else None,
            to_int(row[ix["年报版本年"]].value) if "年报版本年" in ix else None,
            text(row[ix["指标口径"]].value) if "指标口径" in ix else None,
            to_int(row[ix["年报PDF页码"]].value) if "年报PDF页码" in ix else None,
            None, None, None, None, None, None, None,
            text(row[ix["备注"]].value) if "备注" in ix else None,
            None,  # evidence_note（中文表无此列）
            official,
            official_check,
            None,  # jcr_release_year：中文表无 JCR 发布年列（前端中文 Tab 不显示）
            None,  # partition_source：中文不按中科院口径
            "official",  # jif_source：国内年报（official）
            None, None,
            src_label, src_url,
            imported_at,
        ))
        stats["official_total"] += 1
        stats["official_filled"] += 1 if official else 0
        stats["official_check_filled"] += 1 if official_check else 0
        if not official:
            stats["official_missing"].append(("zh", name))
        if composite is None:
            stats["cif_missing"].append((name, check_status or ""))
        stats["hl_zh"] += 1 if src_url else 0
        stats["id_slug"] += 1
    return rows


# ---- 共享主流程 ----

def extract_rows_from_xlsx(
    source: Path,
    *,
    imported_at: str | None = None,
    stats: dict | None = None,
) -> tuple[list[tuple], list[tuple], dict]:
    """xlsx → (rows_en, rows_zh, stats)。口径不符/id 冲突抛 JournalSeedError（不写库）。"""
    stats = stats if stats is not None else new_stats()
    imported_at = imported_at or datetime.now().isoformat(timespec="seconds")

    try:
        wb = openpyxl.load_workbook(source, data_only=True)
    except FileNotFoundError as exc:
        raise JournalSeedError(
            f"源文件不存在: {source}", exit_code=3, stage="读取失败"
        ) from exc
    except Exception as exc:  # noqa: BLE001
        raise JournalSeedError(
            f"无法打开 Excel: {exc}", exit_code=3, stage="读取失败"
        ) from exc
    if SHEET_INTL not in wb.sheetnames or SHEET_ZH not in wb.sheetnames:
        raise JournalSeedError(
            f"缺少 sheet：需要 {SHEET_INTL}/{SHEET_ZH}，实际 {wb.sheetnames}",
            exit_code=3, stage="结构错误",
        )

    rows_en = read_intl(wb[SHEET_INTL], imported_at, stats)
    rows_zh = read_zh(wb[SHEET_ZH], imported_at, stats)
    stats["read_intl"], stats["read_zh"] = len(rows_en), len(rows_zh)
    for r in rows_en:
        stats["part_dist"][r[4] or ("未检得" if r[5] else "空白")] += 1

    # 口径行数校验（不符 → 抛错，不写库）
    if len(rows_en) != EXPECT_INTL or len(rows_zh) != EXPECT_ZH:
        raise JournalSeedError(
            f"行数与口径不符：国际 {len(rows_en)}/{EXPECT_INTL}、"
            f"中文 {len(rows_zh)}/{EXPECT_ZH}，中止导入",
            exit_code=2, stage="行数校验失败",
        )

    # id 冲突检测（同 lang 内重复 → 抛错，不写库）
    ids = [r[0] for r in rows_en + rows_zh]
    dup = [i for i, c in Counter(ids).items() if c > 1]
    if dup:
        conflicts = [f"{i} × {Counter(ids)[i]}" for i in sorted(dup)]
        raise JournalSeedError(
            f"发现 {len(conflicts)} 个重复 id，中止导入（未写库）",
            exit_code=1, stage="id 冲突中止", conflicts=conflicts,
        )

    return rows_en, rows_zh, stats


def upsert_to_db(db_path: Path, rows_en: list[tuple], rows_zh: list[tuple]) -> dict[str, int]:
    """建表(IF NOT EXISTS)+幂等加列+upsert → {"en","zh","total"}。"""
    db_path = Path(db_path)
    db_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(str(db_path))
    try:
        conn.executescript(SCHEMA_SQL)
        migrate(conn)
        conn.executemany(UPSERT_SQL, list(rows_en) + list(rows_zh))
        conn.commit()
        d = dict(conn.execute(
            "SELECT lang, COUNT(*) FROM journal_directory GROUP BY lang").fetchall())
        return {"en": d.get("en", 0), "zh": d.get("zh", 0), "total": sum(d.values())}
    finally:
        conn.close()


def ensure_journal_directory_seeded(
    db_path: Path | None = None,
    source_xlsx: Path | None = None,
) -> dict:
    """启动自动 seed：空表才导入；任何失败收敛为 dict，绝不抛异常。

    返回统一四键：seeded/reason/error/rows（另有 counts 仅成功时给出 en/zh 明细）。
    - 已有数据（只认空/非空，不做补差）→ seeded=False, reason="already has N rows, skip"
    - 空表导入成功 → seeded=True, rows=1056
    - xlsx 缺失/结构错误/口径不符/id 冲突 → seeded=False, error="..."
    """
    settings = get_settings()
    db = Path(db_path) if db_path is not None else settings.db_path_abs
    source = Path(source_xlsx) if source_xlsx is not None else settings.journal_directory_seed_abs

    try:
        if db.exists():
            conn = sqlite3.connect(str(db))
            try:
                row = conn.execute("SELECT COUNT(*) FROM journal_directory").fetchone()
            except sqlite3.OperationalError:
                row = None  # 表不存在 → 视为空
            finally:
                conn.close()
        else:
            row = None
        if row is not None and int(row[0]) > 0:
            n = int(row[0])
            return {
                "seeded": False, "reason": f"already has {n} rows, skip",
                "error": None, "rows": n,
            }

        rows_en, rows_zh, _ = extract_rows_from_xlsx(source)
        counts = upsert_to_db(db, rows_en, rows_zh)
        return {
            "seeded": True, "reason": None, "error": None,
            "rows": counts["total"], "counts": counts,
        }
    except JournalSeedError as exc:
        return {"seeded": False, "reason": None, "error": str(exc), "rows": None}
    except Exception as exc:  # noqa: BLE001
        return {"seeded": False, "reason": None, "error": str(exc), "rows": None}
