"""Synthetic documents for the real stable-DSH conversation acceptance test."""
import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches as WordInches
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt


def main():
    mode = sys.argv[1]
    output = Path.cwd() / "documents"
    output.mkdir(exist_ok=True)
    if mode == "create":
        # These measured values are a fixture, not a claimed biological result.
        image = Image.new("RGB", (800, 400), "white")
        draw = ImageDraw.Draw(image)
        draw.rectangle((100, 180, 280, 350), fill="#168b91")
        draw.rectangle((480, 80, 660, 350), fill="#695bcc")
        image.save(output / "fixture.png")
        document = Document()
        for style in ["Normal", "Title", "Heading 1"]:
            document.styles[style].font.name = "Microsoft YaHei"
            document.styles[style]._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        document.add_heading("中文科研报告 / Research report", 0)
        document.add_paragraph("这是合成测试数据，用于验证文档生成。")
        document.add_heading("结果 / Results", 1)
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "分组 / Group"
        table.rows[0].cells[1].text = "数值 / Value"
        for name, value in [("Control", "12"), ("Treatment", "18")]:
            cells = table.add_row().cells
            cells[0].text, cells[1].text = name, value
        document.add_picture(str(output / "fixture.png"), width=WordInches(5.6))
        document.add_paragraph("Figure 1. Synthetic values: 12 and 18.")
        document.save(output / "report-v1.docx")

        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        title = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12), Inches(1))
        title.text = "中文科研汇报 / Research presentation"
        for run in title.text_frame.paragraphs[0].runs:
            run.font.name, run.font.size = "Microsoft YaHei", Pt(28)
        slide.shapes.add_picture(str(output / "fixture.png"), Inches(0.8), Inches(2), width=Inches(5.6))
        data = CategoryChartData()
        data.categories = ["Control", "Treatment"]
        data.add_series("Synthetic measurement", [12, 18])
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7), Inches(2), Inches(5.5), Inches(4), data)
        slide.notes_slide.notes_text_frame.text = "Synthetic fixture; preserve source values 12 and 18."
        deck.save(output / "slides-v1.pptx")
    elif mode == "edit":
        document = Document(output / "report-v1.docx")
        document.paragraphs[0].runs[0].text = "修改后的科研报告 / Revised research report"
        document.add_paragraph("Revision 2: keep both original values, the table and the embedded figure.")
        document.save(output / "report-v2.docx")
        deck = Presentation(output / "slides-v1.pptx")
        deck.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].text = "修改后的汇报 / Revised presentation"
        deck.save(output / "slides-v2.pptx")
    elif mode == "verify":
        from pypdf import PdfReader
        document = Document(output / "report-v2.docx")
        assert "Revised research report" in document.paragraphs[0].text
        assert len(document.inline_shapes) == 1
        assert [[c.text for c in r.cells] for r in document.tables[0].rows][1:] == [["Control", "12"], ["Treatment", "18"]]
        deck = Presentation(output / "slides-v2.pptx")
        assert len(deck.slides) == 1
        assert "Revised presentation" in deck.slides[0].shapes[0].text
        chart = next(s.chart for s in deck.slides[0].shapes if s.has_chart)
        assert tuple(chart.series[0].values) == (12.0, 18.0)
        assert any(s.shape_type == 13 for s in deck.slides[0].shapes)
        for name in ["report-v1", "report-v2", "slides-v1", "slides-v2"]:
            pdf = PdfReader(output / (name + ".pdf"))
            text = "".join(p.extract_text() or "" for p in pdf.pages)
            assert pdf.pages and ("Research" in text or "Revised" in text), (name, text)
            assert "12" in text and "18" in text, (name, text)
            assert any("\u4e00" <= c <= "\u9fff" for c in text), (name, "Chinese text missing")
        print(json.dumps({"editable_documents": 4, "pdfs": 4, "chinese": True, "embedded_images": True, "native_chart": True}))
    else:
        raise ValueError(mode)


if __name__ == "__main__":
    main()
