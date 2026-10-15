"""Gene-family analysis scope resolution.

The legacy pipeline uses ``prefix`` as both an output namespace and an
analysis-range identifier.  This module keeps ``prefix`` for compatibility and
adds an explicit scope object that records family, species roles, selected
species and resolved input files.
"""

from __future__ import annotations

from dataclasses import asdict, dataclass, field
from pathlib import Path
import hashlib
import json
import os
import re
from typing import Any, Iterable


ROLE_TARGET = "target"
ROLE_REFERENCE = "reference"
ROLE_COMPARATIVE = "comparative"
ROLE_SEED_REFERENCE = "seed_reference"


@dataclass
class ScopeFile:
    original_file_name: str
    species: str
    file_type: str
    role: str
    file_scope: str = ""
    gene_family: str = ""
    ref: bool = False
    reference_use: str = ""
    path: str = ""
    exists: bool = False
    size: int | None = None
    mtime: float | None = None


@dataclass
class GeneFamilyScope:
    family: str
    prefix: str
    scope_id: str
    scope_mode: str
    target_species: list[str] = field(default_factory=list)
    reference_species: list[str] = field(default_factory=list)
    primary_reference_species: list[str] = field(default_factory=list)
    identify_blast_reference_species: list[str] = field(default_factory=list)
    comparative_species: list[str] = field(default_factory=list)
    selected_species: list[str] = field(default_factory=list)
    config_path: str = ""
    files: dict[str, dict[str, list[dict[str, Any]]]] = field(default_factory=dict)
    analysis: dict[str, str] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.lower() in {"none", "nan", "null"}:
        return ""
    return text


def split_csv(value: Any) -> list[str]:
    text = clean_text(value)
    if not text:
        return []
    items = re.split(r"[,;，、\s]+", text)
    return unique_keep_order(item.strip() for item in items if item.strip())


def unique_keep_order(items: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    ordered: list[str] = []
    for item in items:
        if item and item not in seen:
            seen.add(item)
            ordered.append(item)
    return ordered


def parse_bool(value: Any) -> bool:
    text = clean_text(value).lower()
    return text in {"1", "true", "ture", "yes", "y", "ref"}


def sanitize_token(value: str) -> str:
    token = re.sub(r"[^A-Za-z0-9_.-]+", "-", clean_text(value))
    token = token.strip(".-_")
    return token or "NA"



def normalize_reference_use(value: Any) -> str:
    text = clean_text(value).lower().replace("-", "_").replace(" ", "_")
    aliases = {
        "blast": "identify_blast",
        "identify": "identify_blast",
        "identify_reference": "identify_blast",
        "identify_blast_reference": "identify_blast",
        "homology_search": "identify_blast",
        "primary": "downstream_primary",
        "main_reference": "downstream_primary",
        "primary_reference": "downstream_primary",
        "downstream": "downstream_primary",
        "downstream_reference": "downstream_primary",
        "all": "both",
        "both_reference": "both",
        "no": "none",
        "false": "none",
    }
    return aliases.get(text, text)

def reference_use_includes(value: Any, purpose: str) -> bool:
    normalized = normalize_reference_use(value)
    if normalized == "both":
        return True
    if purpose == "identify_blast":
        return normalized in {"identify_blast"}
    if purpose == "downstream_primary":
        return normalized in {"downstream_primary"}
    return False

def normalize_role(role: Any, ref: Any = None, file_scope: Any = None) -> str:
    text = clean_text(role).lower().replace("-", "_")
    scope = clean_text(file_scope).lower().replace("-", "_")
    if text in {"target", "targets"}:
        return ROLE_TARGET
    if text in {"comparative", "comparison", "compare", "outgroup"}:
        return ROLE_COMPARATIVE
    if text in {"seed_reference", "seed", "known_reference"}:
        return ROLE_SEED_REFERENCE
    if text in {"ref", "reference", "references"}:
        if scope == "known_family":
            return ROLE_SEED_REFERENCE
        return ROLE_REFERENCE
    if parse_bool(ref):
        if scope == "known_family":
            return ROLE_SEED_REFERENCE
        return ROLE_REFERENCE
    return ""


def normalize_file_type(file_type: Any, file_scope: Any = None) -> str:
    text = clean_text(file_type).lower().replace("-", "_")
    scope = clean_text(file_scope).lower().replace("-", "_")
    aliases = {
        "protein": "pep",
        "all_pep": "pep",
        "proteome": "pep",
        "faa": "pep",
        "fa_pep": "pep",
        "fasta_pep": "pep",
        "annotation": "gff",
        "gff3": "gff",
        "gtf": "gff",
        "dna": "genome",
        "genomic": "genome",
        "cdna": "cds",
    }
    text = aliases.get(text, text)
    if text == "pep" and scope == "known_family":
        return "family_pep"
    if text == "cds" and scope == "known_family":
        return "family_cds"
    return text


def find_config_file(input_dir: Path) -> Path | None:
    candidates = [input_dir / "config.xlsx", input_dir / "config.xlsm", input_dir / "config.xls"]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    matches = sorted(input_dir.glob("*config*.xlsx"))
    return matches[0] if matches else None


def read_config_workbook(config_path: Path) -> tuple[list[dict[str, Any]], dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except Exception as exc:  # pragma: no cover - dependency issue is environment-specific.
        raise RuntimeError("openpyxl is required to read config.xlsx") from exc

    wb = load_workbook(config_path, data_only=True)
    rows: list[dict[str, Any]] = []
    analysis: dict[str, str] = {}

    if "Analysis" in wb.sheetnames:
        ws = wb["Analysis"]
        headers = [clean_text(cell.value) for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        key_idx = headers.index("Key") if "Key" in headers else 0
        value_idx = headers.index("Value") if "Value" in headers else 1
        for row in ws.iter_rows(min_row=2, values_only=True):
            if len(row) <= key_idx:
                continue
            key = clean_text(row[key_idx])
            if not key:
                continue
            value = clean_text(row[value_idx] if len(row) > value_idx else "")
            analysis[key] = value

    sheet_name = "config" if "config" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]
    header_values = [clean_text(cell.value) for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    for raw in ws.iter_rows(min_row=2, values_only=True):
        row = {header_values[idx]: raw[idx] if idx < len(raw) else None for idx in range(len(header_values))}
        original_name = Path(clean_text(row.get("Original_File_Name") or "")).name
        if original_name.startswith("Ref_combined."):
            continue
        if any(clean_text(value) for value in row.values()):
            rows.append(row)

    return rows, analysis


def resolve_existing_path(input_dir: Path, workspace: Path | None, species: str, family: str, row: dict[str, Any], file_type: str) -> Path | None:
    original = clean_text(row.get("Original_File_Name") or row.get("File_Name") or row.get("filename"))
    names: list[str] = []
    if original:
        names.extend([
            original,
            Path(original).name,
            str(Path(species) / Path(original).name),
        ])

    if species:
        if file_type == "pep":
            names.extend([f"{species}.pep.fasta", f"{species}.pep.fa", f"{species}.protein.fasta"])
        elif file_type == "family_pep":
            names.extend([
                f"{species}.{family}.pep.fa",
                f"{species}.{family}.pep.fasta",
                str(Path(species) / f"{species}.{family}.pep.fasta"),
                str(Path(species) / f"{species}.{family}.pep.fa"),
            ])
        elif file_type == "cds":
            names.extend([f"{species}.cds.fasta", f"{species}.cds.fa"])
        elif file_type == "family_cds":
            names.extend([f"{species}.{family}.cds.fa", f"{species}.{family}.cds.fasta"])
        elif file_type == "gff":
            names.extend([f"{species}_final.gff3", f"{species}.gff3", f"{species}.gff"])
        elif file_type == "genome":
            names.extend([f"{species}.genome.fasta", f"{species}.genome.fa", f"{species}.fasta"])

    search_dirs = [input_dir]
    if workspace:
        search_dirs.extend([workspace / "final", workspace])
    for base in search_dirs:
        for name in unique_keep_order(names):
            candidate = base / name
            if candidate.exists():
                return candidate.resolve()

    for base in reference_search_roots(input_dir, workspace):
        for name in unique_keep_order(names):
            filename = Path(name).name
            if not filename:
                continue
            candidate = base / filename
            if candidate.exists():
                return candidate.resolve()
            try:
                matches = sorted(path for path in base.rglob(filename) if path.is_file())
            except Exception:
                matches = []
            if matches:
                return matches[0].resolve()
    return None


def reference_search_roots(input_dir: Path, workspace: Path | None = None) -> list[Path]:
    """Local reference-data roots used to fill config paths without user prompts."""
    candidates: list[Path] = []
    for base in [input_dir, *list(input_dir.parents)[:4]]:
        candidates.extend([
            base / "ref_data",
            base / "reference",
            base / "references",
            base / "data" / "ref_data",
        ])
    if workspace:
        for base in [workspace, *list(workspace.parents)[:4]]:
            candidates.extend([
                base / "ref_data",
                base / "reference",
                base / "references",
                base / "data" / "ref_data",
            ])
    configured_ref = os.environ.get("PLANTAI_REF_DATA_ROOT")
    if configured_ref:
        candidates.append(Path(configured_ref).expanduser())

    roots: list[Path] = []
    seen: set[str] = set()
    for candidate in candidates:
        try:
            resolved = candidate.resolve()
        except Exception:
            continue
        key = str(resolved)
        if key in seen or not resolved.exists() or not resolved.is_dir():
            continue
        seen.add(key)
        roots.append(resolved)
    return roots


def file_record(path: Path | None) -> tuple[str, bool, int | None, float | None]:
    if not path or not path.exists():
        return "", False, None, None
    stat = path.stat()
    return str(path), True, stat.st_size, stat.st_mtime


def infer_family(prefix: str, analysis: dict[str, str], rows: list[dict[str, Any]], family_override: str = "") -> str:
    if clean_text(family_override):
        return clean_text(family_override)
    for key in ("Family", "Gene_Family", "family", "gene_family"):
        if clean_text(analysis.get(key)):
            return clean_text(analysis.get(key))
    if "." in clean_text(prefix):
        return clean_text(prefix).split(".", 1)[1]
    for row in rows:
        family = clean_text(row.get("Gene_Family"))
        if family:
            return family
    return ""


def default_prefix(prefix: str, family: str, target_species: list[str]) -> str:
    provided = clean_text(prefix)
    if provided:
        return provided
    if family and len(target_species) == 1:
        return f"{target_species[0]}.{family}"
    if family:
        return f"Target.{family}"
    return "Target"


def build_scope_id(family: str, targets: list[str], refs: list[str], comps: list[str], mode: str) -> str:
    parts = [sanitize_token(family or "family")]
    if targets:
        parts.append("target-" + "-".join(sanitize_token(x) for x in targets))
    if refs:
        parts.append("ref-" + "-".join(sanitize_token(x) for x in refs))
    if comps:
        parts.append("comp-" + "-".join(sanitize_token(x) for x in comps))
    if mode and mode != "custom":
        parts.append(sanitize_token(mode))
    return ".".join(parts)


def scope_hash(scope_id: str, length: int = 8) -> str:
    text = sanitize_token(scope_id)
    return hashlib.sha1(text.encode("utf-8")).hexdigest()[: max(4, int(length or 8))]


def summarize_species_list(role: str, species: list[str], max_items: int = 3) -> str:
    clean = [sanitize_token(x) for x in species if clean_text(x)]
    if not clean:
        return ""
    head = clean[:max_items]
    suffix = f"-plus{len(clean) - max_items}" if len(clean) > max_items else ""
    return f"{sanitize_token(role)}-" + "-".join(head) + suffix


def build_scope_summary(
    family: str,
    targets: list[str],
    refs: list[str],
    comps: list[str],
    mode: str,
    max_items: int = 3,
) -> str:
    """Human-readable, bounded scope label for filenames/UI.

    The full machine scope remains in ``scope_id`` and scope JSON.  This
    summary intentionally avoids listing every species when users analyze many
    species.
    """
    parts: list[str] = []
    for role, values in (("target", targets), ("ref", refs), ("comp", comps)):
        part = summarize_species_list(role, values, max_items=max_items)
        if part:
            parts.append(part)
    if not parts:
        parts.append(sanitize_token(family or "family"))
    if mode and mode != "custom":
        parts.append(sanitize_token(mode))
    return ".".join(parts)


def selected_species_for_mode(targets: list[str], refs: list[str], comps: list[str], mode: str) -> list[str]:
    mode = clean_text(mode) or "custom"
    if mode == "target_only":
        return unique_keep_order(targets)
    if mode == "target_plus_reference":
        return unique_keep_order([*targets, *refs])
    if mode == "target_plus_comparative":
        return unique_keep_order([*targets, *comps])
    if mode == "target_plus_reference_plus_comparative":
        return unique_keep_order([*targets, *refs, *comps])
    return unique_keep_order([*targets, *refs, *comps])


def build_gene_family_scope(
    input_dir: str | Path,
    workspace: str | Path | None = None,
    prefix: str = "",
    family: str = "",
    target_species: Iterable[str] | str | None = None,
    reference_species: Iterable[str] | str | None = None,
    comparative_species: Iterable[str] | str | None = None,
    scope_mode: str = "custom",
) -> GeneFamilyScope:
    input_path = Path(input_dir).resolve()
    workspace_path = Path(workspace).resolve() if workspace else None
    config_path = find_config_file(input_path)
    rows: list[dict[str, Any]] = []
    analysis: dict[str, str] = {}
    warnings: list[str] = []

    if config_path:
        rows, analysis = read_config_workbook(config_path)
    else:
        warnings.append(f"No config.xlsx found in {input_path}")

    family_value = infer_family(prefix, analysis, rows, family)

    role_species: dict[str, list[str]] = {
        ROLE_TARGET: [],
        ROLE_REFERENCE: [],
        ROLE_COMPARATIVE: [],
        ROLE_SEED_REFERENCE: [],
    }
    scope_files: dict[str, dict[str, list[dict[str, Any]]]] = {}

    for row in rows:
        species = clean_text(row.get("Species_Code") or row.get("Species") or row.get("species"))
        if not species:
            continue
        file_scope = clean_text(row.get("File_Scope"))
        role = normalize_role(row.get("Species_Role"), row.get("Ref"), file_scope)
        reference_use = normalize_reference_use(row.get("Reference_Use"))
        if role:
            role_species.setdefault(role, [])
            if species not in role_species[role]:
                role_species[role].append(species)

        file_type = normalize_file_type(row.get("File_Type"), file_scope)
        if not file_type:
            continue
        row_family = clean_text(row.get("Gene_Family"))
        if row_family and family_value and row_family.lower() != family_value.lower() and file_type in {"family_pep", "family_cds"}:
            continue
        path = resolve_existing_path(input_path, workspace_path, species, family_value, row, file_type)
        path_text, exists, size, mtime = file_record(path)
        rec = ScopeFile(
            original_file_name=clean_text(row.get("Original_File_Name")),
            species=species,
            file_type=file_type,
            role=role,
            file_scope=file_scope,
            gene_family=row_family,
            ref=parse_bool(row.get("Ref")),
            reference_use=reference_use,
            path=path_text,
            exists=exists,
            size=size,
            mtime=mtime,
        )
        scope_files.setdefault(species, {}).setdefault(file_type, []).append(asdict(rec))

    analysis_targets = split_csv(analysis.get("Target_Species"))
    analysis_refs = split_csv(analysis.get("Reference_Species"))
    analysis_primary_refs = split_csv(analysis.get("Primary_Reference_Species"))
    analysis_identify_blast_refs = split_csv(analysis.get("Identify_Blast_Reference_Species"))
    analysis_comps = split_csv(analysis.get("Comparative_Species"))

    target_provided = target_species is not None
    ref_provided = reference_species is not None
    comp_provided = comparative_species is not None

    target_list = split_csv(target_species) if isinstance(target_species, str) else unique_keep_order(target_species or [])
    ref_list = split_csv(reference_species) if isinstance(reference_species, str) else unique_keep_order(reference_species or [])
    comp_list = split_csv(comparative_species) if isinstance(comparative_species, str) else unique_keep_order(comparative_species or [])

    if not target_provided and not target_list:
        target_list = analysis_targets or role_species.get(ROLE_TARGET, [])
    explicit_primary_refs = analysis_primary_refs or analysis_refs
    if not ref_provided and not ref_list:
        row_primary_refs = unique_keep_order(
            row.get("Species_Code") or row.get("Species") or row.get("species")
            for row in rows
            if reference_use_includes(row.get("Reference_Use"), "downstream_primary")
        )
        ref_list = explicit_primary_refs or row_primary_refs or role_species.get(ROLE_REFERENCE, [])
    identify_blast_ref_list = analysis_identify_blast_refs
    if not identify_blast_ref_list:
        identify_blast_ref_list = unique_keep_order(
            row.get("Species_Code") or row.get("Species") or row.get("species")
            for row in rows
            if reference_use_includes(row.get("Reference_Use"), "identify_blast")
        )
    if not identify_blast_ref_list:
        identify_blast_ref_list = unique_keep_order([*role_species.get(ROLE_SEED_REFERENCE, [])])
    if analysis_identify_blast_refs and ref_list:
        extra = [sp for sp in analysis_identify_blast_refs if sp not in ref_list]
        if extra:
            warnings.append("Identify_Blast_Reference_Species is used only for identify BLAST and will not be added to downstream reference_species: " + ",".join(extra))
    if not analysis_primary_refs and analysis_refs:
        warnings.append("Using legacy Analysis.Reference_Species as downstream primary reference; prefer Primary_Reference_Species for new projects.")
    if not comp_provided and not comp_list:
        comp_list = analysis_comps or role_species.get(ROLE_COMPARATIVE, [])

    prefix_value = default_prefix(prefix, family_value, target_list)
    selected = selected_species_for_mode(target_list, ref_list, comp_list, scope_mode)
    scope_id = build_scope_id(family_value, target_list, ref_list, comp_list, clean_text(scope_mode) or "custom")

    return GeneFamilyScope(
        family=family_value,
        prefix=prefix_value,
        scope_id=scope_id,
        scope_mode=clean_text(scope_mode) or "custom",
        target_species=target_list,
        reference_species=ref_list,
        primary_reference_species=ref_list,
        identify_blast_reference_species=identify_blast_ref_list,
        comparative_species=comp_list,
        selected_species=selected,
        config_path=str(config_path.resolve()) if config_path else "",
        files=scope_files,
        analysis=analysis,
        warnings=warnings,
    )


def write_scope(scope: GeneFamilyScope, output_path: str | Path) -> Path:
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(scope.to_dict(), indent=2, ensure_ascii=False), encoding="utf-8")
    return path


def read_scope(scope_path: str | Path) -> dict[str, Any]:
    return json.loads(Path(scope_path).read_text(encoding="utf-8"))


def scope_manifest_path(workspace: str | Path, prefix: str, scope_id: str) -> Path:
    return Path(workspace) / "final" / f"{sanitize_token(prefix)}.{sanitize_token(scope_id)}.scope.json"


def sha256_file(path: str | Path, chunk_size: int = 1024 * 1024) -> str:
    hasher = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(chunk_size), b""):
            hasher.update(chunk)
    return hasher.hexdigest()
