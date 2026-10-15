#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import json
import re
import time
from pathlib import Path
from typing import Any
import sys

import yaml

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from cli.autofill_common import FASTQ_SUFFIX_RE, backup_if_exists, parse_key_values, scan_files, write_json
from gfa_core.rnaseq_de_design import infer_rnaseq_de_design


SPECIES_DEFAULTS = {
    "Os": {"speciesName": "Oryza sativa", "keggOrganism": "osa"},
    "osa": {"speciesName": "Oryza sativa", "keggOrganism": "osa"},
    "osa1": {"speciesName": "Oryza sativa", "keggOrganism": "osa"},
    "osa1_r7": {"speciesName": "Oryza sativa", "keggOrganism": "osa"},
    "At": {"speciesName": "Arabidopsis thaliana", "keggOrganism": "ath"},
    "Ath": {"speciesName": "Arabidopsis thaliana", "keggOrganism": "ath"},
    "Zm": {"speciesName": "Zea mays", "keggOrganism": "zma"},
    "Gm": {"speciesName": "Glycine max", "keggOrganism": "gmx"},
    "Sl": {"speciesName": "Solanum lycopersicum", "keggOrganism": "sly"},
    "Cp": {"speciesName": "Carica papaya", "keggOrganism": "cpap"},
}


TREATMENT_KEYWORD_DEFAULTS = {
    "salt stress": ["salt stress", "salinity", "NaCl", "盐胁迫", "盐处理"],
    "drought stress": ["drought stress", "water deficit", "dehydration", "干旱胁迫", "干旱"],
    "heat stress": ["heat stress", "high temperature", "热胁迫", "高温胁迫"],
    "cold stress": ["cold stress", "low temperature", "冷胁迫", "低温胁迫"],
    "cadmium stress": ["cadmium stress", "Cd stress", "heavy metal", "镉胁迫", "镉处理"],
    "pathogen stress": ["pathogen stress", "infection", "disease resistance", "病原", "侵染"],
}


CONFIRMED_DE_WARNING_MARKERS = (
    "DE design pending user input",
    "缺少 group_role",
    "唯一对照组",
    "control/compare/ignore",
)


def clean_confirmed_autofill_warnings(warnings: list[Any]) -> list[Any]:
    """Drop draft-only DE design warnings after samples.yaml confirmation succeeds."""
    cleaned = []
    for warning in warnings or []:
        text = str(warning or "")
        if any(marker in text for marker in CONFIRMED_DE_WARNING_MARKERS):
            continue
        cleaned.append(warning)
    return cleaned


def treatment_keywords_from_text(treatment: str) -> list[str]:
    text = str(treatment or "").strip()
    if not text:
        return []
    lower = text.lower()
    keywords = [text]
    for label, values in TREATMENT_KEYWORD_DEFAULTS.items():
        markers = [label, *values]
        if any(marker.lower() in lower for marker in markers):
            keywords.extend(values)
            break
    seen = set()
    out = []
    for keyword in keywords:
        clean = str(keyword or "").strip()
        key = clean.lower()
        if clean and key not in seen:
            seen.add(key)
            out.append(clean)
    return out


def strip_fastq_suffix(name: str) -> str:
    return FASTQ_SUFFIX_RE.sub("", name)


def infer_read_pair(path: Path) -> tuple[str, str]:
    stem = strip_fastq_suffix(path.name)
    patterns = [
        (r"(.+?)_1\.clean$", "1"),
        (r"(.+?)_2\.clean$", "2"),
        (r"(.+?)(?:[_\.-])R?1(?:[_\.-].*)?$", "1"),
        (r"(.+?)(?:[_\.-])R?2(?:[_\.-].*)?$", "2"),
    ]
    for pattern, read in patterns:
        match = re.match(pattern, stem, re.I)
        if match:
            sample = re.sub(r"(^.+rnaseq_qc\.|\.clean$)", "", match.group(1))
            return sample, read
    return stem, ""


def infer_sample_meta(sample_id: str) -> dict[str, Any]:
    clean = re.sub(r"^.+rnaseq_qc\.", "", sample_id)
    if re.match(r"^(?:SRR|ERR|DRR)\d+$", clean, re.I):
        return {"condition": "", "time": "", "replicate": ""}
    rep = ""
    match = re.search(r"(?:^|[_\.-])r(?:ep)?(\d+)$", clean, re.I)
    if match:
        rep = match.group(1)
        condition = clean[:match.start()].rstrip("_.-")
    else:
        match = re.search(r"(?:^|[_\.-])(\d+)$", clean)
        if match:
            rep = match.group(1)
            condition = clean[:match.start()].rstrip("_.-")
        else:
            condition = clean
    timepoint = ""
    tmatch = re.search(r"(?:^|[_\.-])T?(\d+)(?:h|d|day)?(?:$|[_\.-])", condition, re.I)
    if tmatch:
        timepoint = tmatch.group(1)
    return {
        "condition": condition or clean,
        "time": int(timepoint) if timepoint.isdigit() else timepoint,
        "replicate": int(rep) if rep.isdigit() else rep,
    }



SAMPLE_INFO_NAME_RE = re.compile(r"(samples?_?info|sample_?sheet|sample_?metadata|sample|metadata|meta|phenotype|group|design|样本|分组|信息)", re.I)
SAMPLE_INFO_SUFFIXES = (".xlsx", ".xlsm", ".xls", ".csv", ".tsv", ".txt")
SAMPLE_EXTRA_FIELDS = (
    "genome_version",
    "genotype",
    "tissue",
    "dose",
    "batch",
    "subject_id",
    "phenotype",
    "library_strandness",
    "note",
)


def normalize_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9\u4e00-\u9fff]+", "", str(value or "").strip().lower())


def clean_cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return int(value)
    text = str(value).strip()
    return "" if text.lower() in {"nan", "none", "null", "undefined"} else value


def canonical_sample_info_key(header: str) -> str:
    key = normalize_header(header)
    raw = str(header or "").strip().lower()
    if raw.startswith("phenotype_") or raw.startswith("trait_"):
        return "phenotype"
    aliases = {
        "run": {"run", "runs", "srr", "err", "drr", "accession", "runaccession", "测序号", "运行号"},
        "sample_id": {"sample", "sampleid", "sample_id", "samplename", "sample_name", "样本", "样本编号", "样本名"},
        "label": {"name", "label", "title", "samplelabel", "sample_label", "样本名称", "名称"},
        "condition": {"condition", "group", "treatment", "处理", "分组", "组别", "实验组", "conditiongroup"},
        "time": {"time", "timepoint", "stage", "hours", "hour", "时间", "时期", "阶段"},
        "replicate": {"replicate", "rep", "biorep", "biologicalreplicate", "重复", "重复编号", "生物学重复"},
        "include": {"include", "use", "selected", "enabled", "是否使用", "纳入"},
        "group_role": {"group_role", "grouprole", "role", "de_role", "derole", "分组角色", "角色"},
        "genome_version": {"genome", "genomeversion", "genome_version", "reference", "referencegenome", "reference_genome", "refgenome", "assembly", "参考基因组"},
        "genotype": {"genotype", "cultivar", "variety", "accession", "line", "strain", "ecotype", "品种", "基因型"},
        "tissue": {"tissue", "organ", "sampletype", "sample_type", "组织", "器官"},
        "dose": {"dose", "dosage", "concentration", "conc", "level", "剂量", "浓度"},
        "batch": {"batch", "librarybatch", "library_batch", "seqbatch", "seq_batch", "lane", "block", "批次", "技术批次"},
        "subject_id": {"subject", "subjectid", "subject_id", "plant", "plantid", "plant_id", "individual", "individualid", "patient", "patientid", "配对", "植株"},
        "phenotype": {"phenotype", "trait", "traits", "表型", "性状"},
        "library_strandness": {"librarystrandness", "library_strandness", "strandness", "strand", "strandedness", "链特异性"},
        "note": {"note", "notes", "comment", "comments", "remark", "remarks", "备注"},
    }
    normalized_aliases = {name: {normalize_header(x) for x in values} for name, values in aliases.items()}
    for name, values in normalized_aliases.items():
        if key in values:
            return name
    return ""


def infer_condition_from_label(label: str, sample_id: str) -> str:
    text = str(label or "").strip()
    if not text:
        return ""
    if text == str(sample_id or "").strip() and re.match(r"^(?:SRR|ERR|DRR)\d+$", text, re.I):
        return ""
    text = re.sub(r"\s+", "_", text)
    text = re.sub(r"(?:[_\.-])(?:rep|r)?\d+$", "", text, flags=re.I)
    text = re.sub(r"(?:[_\.-])\d+$", "", text)
    return text.strip("_.-")


def find_sample_info_path(source_dir: Path, input_dir: Path) -> Path | None:
    candidates: list[Path] = []
    search_roots = []
    for root in [source_dir, source_dir.parent if source_dir else None, input_dir]:
        if root and root.exists() and root not in search_roots:
            search_roots.append(root)
    for root in search_roots:
        iterator = root.rglob("*") if root.is_dir() else [root]
        for path in iterator:
            if not path.is_file():
                continue
            lower = path.name.lower()
            if lower.endswith(SAMPLE_INFO_SUFFIXES) and SAMPLE_INFO_NAME_RE.search(path.name):
                candidates.append(path)
    return sorted(candidates, key=lambda p: (p.parent != input_dir, len(p.parts), len(p.name), p.name))[0] if candidates else None


def read_sample_info_table(path: Path) -> tuple[dict[str, dict[str, Any]], list[str]]:
    warnings: list[str] = []
    if not path or not path.exists():
        return {}, [f"Sample info table not found: {path}"]
    raw_rows: list[dict[str, Any]] = []
    suffix = path.suffix.lower()
    try:
        if suffix in {".xlsx", ".xlsm", ".xls"}:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                return {}, [f"Sample info table is empty: {path}"]
            headers = [str(h or "").strip() for h in rows[0]]
            for values in rows[1:]:
                raw_rows.append({headers[i]: clean_cell(values[i] if i < len(values) else "") for i in range(len(headers))})
        else:
            delimiter = "\t" if suffix in {".tsv", ".txt"} else ","
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                reader = csv.DictReader(handle, delimiter=delimiter)
                raw_rows = [{k: clean_cell(v) for k, v in row.items()} for row in reader]
    except Exception as exc:
        return {}, [f"Could not read sample info table {path}: {exc}"]
    by_sample: dict[str, dict[str, Any]] = {}
    for raw in raw_rows:
        item: dict[str, Any] = {}
        for header, value in raw.items():
            key = canonical_sample_info_key(header)
            if key:
                cleaned = clean_cell(value)
                if key == "phenotype" and cleaned not in ("", None):
                    header_text = str(header or "").strip()
                    item[key] = f"{header_text}={cleaned}" if header_text.lower().startswith(("phenotype_", "trait_")) else cleaned
                else:
                    item[key] = cleaned
        run = str(item.get("run") or "").strip()
        sample_id = str(item.get("sample_id") or "").strip()
        label = str(item.get("label") or "").strip()
        if not run and sample_id and re.match(r"^(?:SRR|ERR|DRR)\d+$", sample_id, re.I):
            run = sample_id
            sample_id = label or sample_id
        if not run and label and re.match(r"^(?:SRR|ERR|DRR)\d+$", label, re.I):
            run = label
        if not run:
            continue
        display_sample_id = sample_id or label or run
        condition = str(item.get("condition") or "").strip()
        if not condition:
            condition = infer_condition_from_label(display_sample_id, run)
        by_sample[run] = {
            "run": run,
            "sample_id": display_sample_id,
            "condition": condition,
            "time": item.get("time", ""),
            "replicate": item.get("replicate", ""),
            "include": item.get("include", True) if item.get("include", "") != "" else True,
            "group_role": item.get("group_role", ""),
            "label": label,
        }
        for field in SAMPLE_EXTRA_FIELDS:
            if item.get(field) not in ("", None):
                by_sample[run][field] = item.get(field)
    if not by_sample:
        warnings.append(f"No usable sample rows found in sample info table: {path}")
    return by_sample, warnings

def find_fastqs(source_dir: Path) -> list[Path]:
    return [p for p in scan_files(source_dir) if FASTQ_SUFFIX_RE.search(p.name)]


def build_samples(source_dir: Path, input_dir: Path, sample_info_path: Path | None = None) -> tuple[dict, list[dict], list[str]]:
    pairs: dict[str, dict[str, Path]] = {}
    warnings: list[str] = []
    sample_info_file = sample_info_path or find_sample_info_path(source_dir, input_dir)
    sample_info, info_warnings = read_sample_info_table(sample_info_file) if sample_info_file else ({}, [])
    warnings.extend(info_warnings)
    for path in find_fastqs(source_dir):
        sample, read = infer_read_pair(path)
        if not read:
            warnings.append(f"Could not infer read pair from {path.name}")
            continue
        pairs.setdefault(sample, {})[f"r{read}"] = path

    samples: dict[str, dict[str, Any]] = {}
    rows: list[dict] = []
    for run, reads in sorted(pairs.items()):
        meta = infer_sample_meta(run)
        table_meta = sample_info.get(run, {})
        sample_id = str(table_meta.get("sample_id") or run).strip()
        if table_meta:
            for key in ("condition", "time", "replicate", "include", "group_role", *SAMPLE_EXTRA_FIELDS):
                if table_meta.get(key) not in ("", None):
                    meta[key] = table_meta.get(key)
        r1 = reads.get("r1")
        r2 = reads.get("r2")
        if not r1 or not r2:
            warnings.append(f"Sample {run} is missing R1 or R2")
            continue
        r1_text = str(r1 if r1.is_absolute() else r1)
        r2_text = str(r2 if r2.is_absolute() else r2)
        sample_record = {
            "run": run,
            "r1": r1_text,
            "r2": r2_text,
            "condition": meta["condition"],
            "time": meta["time"],
            "replicate": meta["replicate"],
            "include": table_meta.get("include", True),
            "group_role": meta.get("group_role", ""),
        }
        for field in SAMPLE_EXTRA_FIELDS:
            if meta.get(field) not in ("", None):
                sample_record[field] = meta.get(field)
        samples[sample_id] = sample_record
        row = {
            "run": run,
            "sample_id": sample_id,
            "fastq_1": r1_text,
            "fastq_2": r2_text,
            "condition": meta["condition"],
            "time": meta["time"],
            "replicate": meta["replicate"],
            "include": table_meta.get("include", True),
            "group_role": meta.get("group_role", ""),
            "confidence": "high" if table_meta and meta["condition"] and meta["replicate"] != "" else "high" if meta["condition"] and meta["replicate"] != "" else "low" if not meta["condition"] else "medium",
        }
        for field in SAMPLE_EXTRA_FIELDS:
            row[field] = meta.get(field, "")
        rows.append(row)
    if sample_info_file:
        matched = sum(1 for run in pairs if run in sample_info)
        missing = sorted(set(pairs) - set(sample_info))
        table_only = [run for run in sample_info if run not in pairs]
        if matched:
            warnings.append(f"Sample info table merged: {sample_info_file} ({matched}/{len(pairs)} FASTQ samples matched)")
        if missing:
            warnings.append(f"Sample info table did not contain {len(missing)} FASTQ sample IDs; first unmatched: {', '.join(missing[:5])}")
        if table_only:
            warnings.append(f"Sample info table contains {len(table_only)} run IDs not found in current FASTQ path; first missing FASTQ runs: {', '.join(table_only[:8])}")
    return samples, rows, warnings


def design_from_samples(samples: dict[str, dict[str, Any]], description: str, design_hint: str, control_hint: str) -> dict:
    hints = parse_key_values(description)
    treatment = str(hints.get("treatment") or "").strip()
    treatment_keywords = treatment_keywords_from_text(treatment)
    plan = infer_rnaseq_de_design(samples)
    design_hint_clean = str(design_hint or "").strip()
    if design_hint_clean.lower() == "manual":
        design_hint_clean = ""
    if design_hint_clean and design_hint_clean != plan["design"]:
        raise ValueError(f"样本表分组角色推断为 {plan['design']}，与 --design={design_hint} 不一致")
    if control_hint and control_hint.strip() != plan["control_group"]:
        raise ValueError(f"样本表 control={plan['control_group']}，与 --control-group={control_hint} 不一致")
    experiment = dict(plan["experiment"])
    if treatment:
        experiment["treatment"] = treatment
    if treatment_keywords:
        experiment["treatment_keywords"] = treatment_keywords
    return {
        "experiment": experiment,
        "compare": plan["compare"],
        "de": plan["de"],
        "groups": plan["groups"],
    }


def parse_time_sort_value(value: Any) -> tuple[int, Any]:
    text = str(value)
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if match:
        number = float(match.group(0))
        return (0, number)
    return (1, text)


def dedupe_pairs(pairs: list[list[str]]) -> list[list[str]]:
    seen = set()
    out = []
    for a, b in pairs:
        key = (a, b)
        if a and b and a != b and key not in seen:
            seen.add(key)
            out.append([a, b])
    return out


def write_yaml(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(yaml.safe_dump(payload, allow_unicode=True, sort_keys=False), encoding="utf-8")


def write_report(path: Path, rows: list[dict], warnings: list[str], design: dict, source_dir: Path) -> None:
    path.write_text(
        "\n".join([
            "# samples.yaml Autofill Report",
            "",
            f"- Source directory: `{source_dir}`",
            f"- Samples: {len(rows)}",
            f"- Design: `{design.get('experiment', {}).get('design', '')}`",
            f"- Control: `{design.get('experiment', {}).get('control_group', '')}`",
            f"- Treatment keyword: `{design.get('experiment', {}).get('treatment', '')}`",
            f"- Pairwise comparisons: {design.get('compare', [])}",
            f"- Warnings: {len(warnings)}",
            "",
            *[f"- {w}" for w in warnings],
        ]) + "\n",
        encoding="utf-8",
    )


def merge_rows(existing: list[dict], detected: list[dict], refresh_metadata: bool = False) -> list[dict]:
    """Merge detected sample rows into an existing draft, preserving user edits unless metadata refresh is explicit."""
    detected_by_id = {}
    for row in detected:
        for field in ("run", "sample_id"):
            key = str(row.get(field) or "").strip()
            if key:
                detected_by_id[key] = row
    merged: list[dict] = []
    seen: set[str] = set()
    metadata_fields = ["run", "sample_id", "condition", "time", "replicate", "include", "confidence", "group_role", *SAMPLE_EXTRA_FIELDS]
    for row in existing:
        key = str(row.get("run") or row.get("sample_id") or "").strip()
        if not key or key in seen:
            continue
        if refresh_metadata and key in detected_by_id:
            updated = dict(row)
            detected_row = detected_by_id[key]
            for field in metadata_fields:
                value = detected_row.get(field)
                if value not in ("", None):
                    updated[field] = value
            for field in ("fastq_1", "fastq_2"):
                if not updated.get(field) and detected_row.get(field):
                    updated[field] = detected_row.get(field)
            row = updated
        merged.append(row)
        seen.add(key)
    for row in detected:
        key = str(row.get("run") or row.get("sample_id") or "").strip()
        if not key or key in seen:
            continue
        merged.append(row)
        seen.add(key)
    return merged


def samples_from_rows(rows: list[dict]) -> dict[str, dict[str, Any]]:
    samples: dict[str, dict[str, Any]] = {}
    for row in rows:
        sample_id = str(row.get("sample_id") or "").strip()
        if not sample_id:
            continue
        record = {
            "run": row.get("run", ""),
            "r1": row.get("fastq_1", ""),
            "r2": row.get("fastq_2", ""),
            "condition": row.get("condition", ""),
            "time": row.get("time", ""),
            "replicate": row.get("replicate", ""),
            "library_strandness": row.get("library_strandness", ""),
            "include": row.get("include", True),
            "group_role": row.get("group_role", ""),
        }
        for field in SAMPLE_EXTRA_FIELDS:
            value = row.get(field, "")
            if value not in ("", None):
                record[field] = value
        samples[sample_id] = record
    return samples


def example_payload() -> dict:
    return {
        "rice_two_group": {
            "note": "Rice treatment-control example with two biological replicates per group. time is blank because this is not a time-series design.",
            "samples": {
                "Os_CK_rep1": {"r1": "Os_CK_rep1_R1.fq.gz", "r2": "Os_CK_rep1_R2.fq.gz", "condition": "CK", "time": "", "replicate": 1},
                "Os_CK_rep2": {"r1": "Os_CK_rep2_R1.fq.gz", "r2": "Os_CK_rep2_R2.fq.gz", "condition": "CK", "time": "", "replicate": 2},
                "Os_salt_rep1": {"r1": "Os_salt_rep1_R1.fq.gz", "r2": "Os_salt_rep1_R2.fq.gz", "condition": "salt", "time": "", "replicate": 1},
                "Os_salt_rep2": {"r1": "Os_salt_rep2_R1.fq.gz", "r2": "Os_salt_rep2_R2.fq.gz", "condition": "salt", "time": "", "replicate": 2},
            },
            "experiment": {"design": "two_group", "control_group": "CK", "treatment_groups": ["salt"], "treatment": "salt stress", "treatment_keywords": ["salt stress", "salinity", "NaCl"]},
            "compare": [["salt", "CK"]],
        },
        "litchi_tissue": {
            "note": "Litchi tissue comparison example with two biological replicates per tissue. Fill time only for time-series or developmental-stage designs.",
            "samples": {
                "SRR25082136": {"r1": "SRR25082136_1.fastq", "r2": "SRR25082136_2.fastq", "condition": "peel", "tissue": "果皮", "time": "", "replicate": 1},
                "SRR25082135": {"r1": "SRR25082135_1.fastq", "r2": "SRR25082135_2.fastq", "condition": "peel", "tissue": "果皮", "time": "", "replicate": 2},
                "SRR25082143": {"r1": "SRR25082143_1.fastq", "r2": "SRR25082143_2.fastq", "condition": "flesh", "tissue": "果肉", "time": "", "replicate": 1},
                "SRR25082141": {"r1": "SRR25082141_1.fastq", "r2": "SRR25082141_2.fastq", "condition": "flesh", "tissue": "果肉", "time": "", "replicate": 2},
            },
            "experiment": {"design": "multi_group", "control_group": "", "treatment_groups": ["peel", "flesh"]},
            "compare": [],
        },
        "time_series": {
            "note": "Use time only when samples are true time points or developmental stages.",
            "samples": {
                "T0_rep1": {"r1": "T0_rep1_R1.fq.gz", "r2": "T0_rep1_R2.fq.gz", "condition": "T0", "time": "0h", "replicate": 1},
                "T0_rep2": {"r1": "T0_rep2_R1.fq.gz", "r2": "T0_rep2_R2.fq.gz", "condition": "T0", "time": "0h", "replicate": 2},
                "T12_rep1": {"r1": "T12_rep1_R1.fq.gz", "r2": "T12_rep1_R2.fq.gz", "condition": "T12", "time": "12h", "replicate": 1},
                "T12_rep2": {"r1": "T12_rep2_R1.fq.gz", "r2": "T12_rep2_R2.fq.gz", "condition": "T12", "time": "12h", "replicate": 2},
            },
            "experiment": {"design": "time_series", "control_group": "T0", "treatment_groups": ["T12"]},
            "compare": [["T12", "T0"]],
            "de": {"strategy": "auto", "design": "time_series", "run_lrt": True},
        },
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", required=True)
    parser.add_argument("--workspace", required=True)
    parser.add_argument("--source-dir", default="")
    parser.add_argument("--sample-info", default="")
    parser.add_argument("--description", default="")
    parser.add_argument("--design", default="")
    parser.add_argument("--control-group", default="")
    parser.add_argument("--species-code", default="")
    parser.add_argument("--species-name", default="")
    parser.add_argument("--kegg-organism", default="")
    parser.add_argument("--confirm", action="store_true")
    parser.add_argument("--draft", default="")
    args = parser.parse_args()

    start = time.time()
    input_dir = Path(args.input_dir)
    workspace = Path(args.workspace)
    final_dir = workspace / "final"
    final_dir.mkdir(parents=True, exist_ok=True)
    source_dir = Path(args.source_dir).expanduser() if args.source_dir else input_dir
    draft_path = Path(args.draft).expanduser() if args.draft else final_dir / "samples.autofill.draft.yaml"
    report_path = final_dir / "samples.autofill.report.md"
    manifest_path = final_dir / "samples_autofill.manifest.json"

    if args.confirm:
        if not draft_path.exists():
            raise SystemExit(f"Draft samples YAML not found: {draft_path}")
        payload = yaml.safe_load(draft_path.read_text(encoding="utf-8")) or {}
        de_warning = ""
        try:
            design = design_from_samples(
                payload.get("samples") if isinstance(payload.get("samples"), dict) else {},
                args.description,
                args.design or str((payload.get("experiment") or {}).get("design") or ""),
                args.control_group or str((payload.get("experiment") or {}).get("control_group") or ""),
            )
        except ValueError as exc:
            design = {"groups": {}, "experiment": {}, "compare": [], "de": {}}
            de_warning = f"DE design pending user input: {exc}"
        payload["groups"] = design["groups"]
        payload["experiment"] = design["experiment"]
        payload["compare"] = design["compare"]
        payload["de"] = design["de"]
        warnings = clean_confirmed_autofill_warnings(payload.get("_autofill", {}).get("warnings", []))
        if de_warning:
            warnings.append(de_warning)
        payload.setdefault("_autofill", {})["warnings"] = warnings
        archive = backup_if_exists(input_dir / "samples.yaml", input_dir / "uploaded_configs", "samples.yaml")
        write_yaml(input_dir / "samples.yaml", payload)
        rows = payload.get("_autofill", {}).get("rows", [])
        warnings = payload.get("_autofill", {}).get("warnings", [])
        status = "success"
        message = "Confirmed draft written to input samples.yaml"
    else:
        sample_info_path = Path(args.sample_info).expanduser() if args.sample_info else None
        detected_samples, detected_rows, warnings = build_samples(source_dir, input_dir, sample_info_path)
        existing_payload = yaml.safe_load(draft_path.read_text(encoding="utf-8")) if draft_path.exists() else {}
        existing_rows = (existing_payload or {}).get("_autofill", {}).get("rows", [])
        rows = merge_rows(existing_rows if isinstance(existing_rows, list) else [], detected_rows, refresh_metadata=bool(sample_info_path or find_sample_info_path(source_dir, input_dir)))
        samples = samples_from_rows(rows)
        if len(rows) > len(detected_rows):
            warnings = list((existing_payload or {}).get("_autofill", {}).get("warnings", [])) + warnings
        try:
            design = design_from_samples(samples, args.description, args.design, args.control_group)
        except ValueError as exc:
            warnings.append(f"DE design pending user input: {exc}")
            design = {"groups": {}, "experiment": {}, "compare": [], "de": {}}
        description_hints = parse_key_values(args.description)
        species_code = str(args.species_code or description_hints.get("species_code") or description_hints.get("target") or "").strip()
        species_defaults = SPECIES_DEFAULTS.get(species_code, {})
        species_name = str(args.species_name or description_hints.get("species_name") or species_defaults.get("speciesName") or "").strip()
        if not species_defaults and species_name:
            for code, defaults in SPECIES_DEFAULTS.items():
                if str(defaults.get("speciesName") or "").lower() == species_name.lower():
                    species_defaults = defaults
                    if not species_code and len(code) <= 3:
                        species_code = code
                    break
        kegg_organism = str(args.kegg_organism or species_defaults.get("keggOrganism") or "").strip().lower()
        payload = {
            "samples": samples,
            **design,
            "replicate_policy": {"min_reps_per_group": 1, "allow_no_replicate": True, "on_violation": "warn"},
            "_autofill": {"rows": rows, "warnings": warnings, "examples": example_payload()},
        }
        if species_code or species_name or kegg_organism:
            payload["species"] = {"code": species_code, "speciesName": species_name, "keggOrganism": kegg_organism}
            payload["speciesCode"] = species_code
            payload["speciesName"] = species_name
            payload["keggOrganism"] = kegg_organism
        write_yaml(draft_path, payload)
        write_report(report_path, rows, warnings, design, source_dir)
        archive = ""
        status = "needs_user_confirmation"
        message = "Draft samples.yaml generated or merged; user confirmation required before replacing samples.yaml"

    manifest = {
        "manifest_version": "1.0",
        "stage": "samples_autofill",
        "status": status,
        "message": message,
        "metrics": {
            "samples": len(rows),
            "warnings": len(warnings),
            "new_detected_samples": len(detected_rows) if not args.confirm else 0,
            "existing_draft_samples": len(existing_rows) if not args.confirm and isinstance(existing_rows, list) else 0,
        },
        "artifacts": {
            "draft_samples": {"path": str(draft_path), "type": "yaml"},
            "report": {"path": str(report_path), "type": "markdown"},
            "confirmed_samples": {"path": str(input_dir / "samples.yaml"), "type": "yaml"} if args.confirm else {},
        },
        "backup": archive,
        "elapsed_sec": round(time.time() - start, 3),
    }
    write_json(manifest_path, manifest)
    print(json.dumps(manifest, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
