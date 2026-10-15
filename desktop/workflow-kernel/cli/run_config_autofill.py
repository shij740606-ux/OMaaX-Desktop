#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
import time
from pathlib import Path
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from cli.autofill_common import (
    CONFIG_COLUMNS,
    CONFIG_REVIEW_COLUMNS,
    backup_if_exists,
    clean_text,
    infer_file_type,
    infer_species_code,
    parse_key_values,
    scan_files,
    write_json,
)
from gfa_core.config_templates import config_template_examples


def _md_escape(value: object) -> str:
    text = clean_text(value)
    return text.replace("|", "\\|").replace("\n", " ")


def _markdown_table(headers: list[str], rows: list[list[object]]) -> list[str]:
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join(["---"] * len(headers)) + " |",
    ]
    for row in rows:
        lines.append("| " + " | ".join(_md_escape(value) for value in row) + " |")
    return lines


def infer_family_from_name(name: str, species: str = "") -> str:
    stem = Path(name).name
    if species:
        prefix = f"{species}."
        if stem.startswith(prefix):
            parts = stem.split(".")
            if len(parts) >= 3 and parts[1].lower() not in {"pep", "cds", "gff", "gff3", "genome"}:
                return parts[1]
    match = re.match(r"^[A-Za-z]{1,6}[._-]([A-Za-z0-9-]+)[._-](?:pep|cds|protein)", stem, re.I)
    if match:
        return match.group(1)
    return ""


def _looks_like_whole_proteome(path: Path) -> bool:
    name = path.name.lower()
    return any(token in name for token in [
        ".pep.all.",
        "pep.all",
        "protein.faa",
        "protein_primary",
        "primarytranscript",
        "primarytranscriptonly",
    ])


def _is_likely_reference_species(path: Path, species: str) -> bool:
    species_text = clean_text(species).lower()
    name = path.name.lower()
    return species_text in {"at", "ath"} or any(token in name for token in ["arabidopsis", "tair"])


def build_rows(input_dir: Path, source_dir: Path, description: str, species_code: str, gene_family: str) -> list[dict]:
    hints = parse_key_values(description)
    species_hint = species_code or hints.get("species_code") or hints.get("target") or ""
    family_hint = gene_family or hints.get("gene_family") or ""
    target_code = species_hint
    rows: list[dict] = []
    for path in scan_files(source_dir):
        if path.parent != source_dir and path.parent.parent == source_dir and path.name.startswith(f"{path.parent.name}."):
            continue
        file_type, type_conf, type_note = infer_file_type(path)
        if not file_type or file_type == "fastq":
            continue
        if path.name == "config.xlsx" or path.suffix.lower() in {".json", ".md", ".txt", ".tsv"}:
            continue
        if path.name.startswith("Ref_combined."):
            continue
        species, species_conf = infer_species_code(path.name, species_hint)
        confidence = "high" if type_conf == "high" and species_conf == "high" else "medium" if type_conf != "low" and species else "low"
        role = "target"
        ref = "FALSE"
        scope = "whole_genome"
        row_family = ""
        reference_use = ""
        lower = path.name.lower()
        inferred_family = infer_family_from_name(path.name, species)
        whole_proteome = file_type == "pep" and _looks_like_whole_proteome(path)
        if family_hint and (not whole_proteome) and any(x in lower for x in [family_hint.lower(), "family", "known", "seed"]):
            row_family = family_hint
            role = "seed_reference"
            ref = "TRUE"
            scope = "known_family"
            reference_use = "identify_blast"
        elif inferred_family and file_type in {"pep", "cds"} and not whole_proteome:
            row_family = inferred_family
            role = "seed_reference"
            ref = "TRUE"
            scope = "known_family"
            reference_use = "identify_blast"
        elif target_code and species and species != target_code and _is_likely_reference_species(path, species):
            role = "reference"
            ref = "TRUE"
            scope = "whole_genome"
            reference_use = "downstream_primary"
        elif target_code and species and species != target_code:
            role = "comparative"
        notes = "; ".join(x for x in [
            type_note,
            "" if species else "Species_Code needs confirmation",
            "Whole-proteome filename kept as whole_genome; only known-family seed rows should use Gene_Family" if whole_proteome else "",
            "Species_Role inferred as reference for Arabidopsis/TAIR whole-genome resource" if role == "reference" else "",
            "Species_Role inferred as comparative because Species_Code differs from target; confirm if this is outgroup or seed_reference/reference" if role == "comparative" else "",
        ] if x)
        if role == "comparative" and confidence == "high":
            confidence = "medium"
        original = path.name if path.parent == input_dir else str(path)
        rows.append({
            "Original_File_Name": original,
            "Species_Code": species,
            "Gene_Family": row_family,
            "File_Type": file_type,
            "Ref": ref,
            "Species_Role": role,
            "File_Scope": scope,
            "Reference_Use": reference_use,
            "confidence": confidence,
            "notes": notes,
        })
    return rows


def write_workbook(path: Path, rows: list[dict], analysis_family: str = "", primary_reference_species: str = "", identify_blast_reference_species: str = "") -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "config"
    ws.append(CONFIG_COLUMNS)
    for row in rows:
        ws.append([row.get(col, "") for col in CONFIG_COLUMNS])
    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    widths = [42, 14, 14, 12, 9, 18, 18, 22]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"

    ex = wb.create_sheet("official_template")
    ex.append(CONFIG_COLUMNS)
    for row in config_template_examples():
        ex.append([row.get(col, "") for col in CONFIG_COLUMNS])
    for cell in ex[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
    for idx, width in enumerate(widths, start=1):
        ex.column_dimensions[get_column_letter(idx)].width = width
    ex.freeze_panes = "A2"

    analysis = wb.create_sheet("Analysis")
    analysis.append(["Key", "Value", "Description"])
    primary_reference_species = clean_text(primary_reference_species)
    identify_blast_reference_species = clean_text(identify_blast_reference_species)
    analysis.append(["Family", clean_text(analysis_family), "Gene family for this job. This is the preferred source when the user does not name a family."])
    analysis.append(["Target_Species", "", "Optional comma-separated target species codes."])
    analysis.append(["Reference_Species", primary_reference_species, "Legacy/compatible downstream reference species. Mirrored from Primary_Reference_Species for old code paths."])
    analysis.append(["Primary_Reference_Species", primary_reference_species, "Main downstream reference species. Usually one species."])
    analysis.append(["Identify_Blast_Reference_Species", identify_blast_reference_species, "Reference family protein species used only as identify BLAST query evidence against the target whole-proteome database; can be multiple."])
    analysis.append(["Comparative_Species", "", "Optional comma-separated comparative/outgroup species codes."])
    analysis.append(["Required_Pfams", "", "Optional Pfam accession(s) for identify."])
    for cell in analysis[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E2F0D9")
        cell.alignment = Alignment(horizontal="center")
    analysis.column_dimensions["A"].width = 24
    analysis.column_dimensions["B"].width = 22
    analysis.column_dimensions["C"].width = 76

    review = wb.create_sheet("_autofill_review")
    review.append(CONFIG_REVIEW_COLUMNS)
    for row in rows:
        review.append([row.get(col, "") for col in CONFIG_REVIEW_COLUMNS])
    for cell in review[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
    review_widths = [42, 14, 14, 12, 9, 18, 18, 22, 12, 70]
    for idx, width in enumerate(review_widths, start=1):
        review.column_dimensions[get_column_letter(idx)].width = width
    review.freeze_panes = "A2"
    wb.save(path)


def read_rows(path: Path) -> list[dict]:
    wb = load_workbook(path)
    ws = wb["config"] if "config" in wb.sheetnames else wb.active
    headers = [clean_text(cell.value) for cell in ws[1]]
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        item = {headers[i]: clean_text(values[i]) if i < len(values) else "" for i in range(len(headers))}
        if clean_text(item.get("Original_File_Name")):
            rows.append(item)
    if "_autofill_review" in wb.sheetnames:
        review_ws = wb["_autofill_review"]
        review_headers = [clean_text(cell.value) for cell in review_ws[1]]
        review_by_name = {}
        for values in review_ws.iter_rows(min_row=2, values_only=True):
            item = {review_headers[i]: clean_text(values[i]) if i < len(values) else "" for i in range(len(review_headers))}
            original = item.get("Original_File_Name")
            if original:
                review_by_name[original] = item
        for row in rows:
            row.update({k: v for k, v in review_by_name.get(row.get("Original_File_Name"), {}).items() if k in {"confidence", "notes"}})
    return rows


def read_analysis_policy(path: Path) -> dict[str, str]:
    policy = {"family": "", "primary_reference_species": "", "identify_blast_reference_species": ""}
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        if "Analysis" not in wb.sheetnames:
            return policy
        ws = wb["Analysis"]
        for values in ws.iter_rows(min_row=2, values_only=True):
            key = clean_text(values[0] if values and len(values) > 0 else "").lower()
            value = clean_text(values[1] if values and len(values) > 1 else "")
            if key in {"family", "gene_family", "analysis_family", "project_family"} and value:
                policy["family"] = value
            elif key == "primary_reference_species" and value:
                policy["primary_reference_species"] = value
            elif key == "identify_blast_reference_species" and value:
                policy["identify_blast_reference_species"] = value
            elif key == "reference_species" and value and not policy["primary_reference_species"]:
                policy["primary_reference_species"] = value
    except Exception:
        return policy
    return policy


def read_analysis_family(path: Path) -> str:
    return read_analysis_policy(path).get("family", "")


def merge_rows(existing: list[dict], detected: list[dict]) -> list[dict]:
    """Merge detected rows into an existing draft without overwriting user edits."""
    merged: list[dict] = []
    seen: set[str] = set()
    for row in existing:
        key = clean_text(row.get("Original_File_Name"))
        if not key or key in seen:
            continue
        merged.append(row)
        seen.add(key)
    for row in detected:
        key = clean_text(row.get("Original_File_Name"))
        if not key or key in seen:
            continue
        merged.append(row)
        seen.add(key)
    return merged


def write_report(path: Path, rows: list[dict], source_dir: Path) -> None:
    low = [r for r in rows if r.get("confidence") == "low"]
    review_rows = [
        [
            row.get("Original_File_Name", ""),
            row.get("Species_Code", ""),
            row.get("File_Type", ""),
            row.get("Species_Role", ""),
            row.get("File_Scope", ""),
            row.get("confidence", ""),
            row.get("notes", ""),
        ]
        for row in rows
    ]
    template_rows = [
        [
            row.get("Original_File_Name", ""),
            row.get("Species_Code", ""),
            row.get("Gene_Family", ""),
            row.get("File_Type", ""),
            row.get("Ref", ""),
            row.get("Species_Role", ""),
            row.get("File_Scope", ""),
            row.get("Reference_Use", ""),
        ]
        for row in config_template_examples()
    ]
    path.write_text(
        "\n".join([
            "# config.xlsx Autofill Report",
            "",
            f"- Source directory: `{source_dir}`",
            f"- Draft rows: {len(rows)}",
            f"- Low-confidence rows: {len(low)}",
            "- Formal Excel sheet: `config`",
            "- Analysis sheet: `Family` stores the job-level gene family; row-level `Gene_Family` is only required for known-family seed files.",
            "- Formal columns: `Original_File_Name`, `Species_Code`, `Gene_Family`, `File_Type`, `Ref`, `Species_Role`, `File_Scope`, `Reference_Use`",
            "- Review-only sheet: `_autofill_review` stores `confidence` and `notes`.",
            "",
            "## Draft Rows",
            "",
            *_markdown_table(["Original file", "Species", "Type", "Role", "Scope", "Confidence", "Notes"], review_rows),
            "",
            "## Official Template Examples",
            "",
            *_markdown_table(CONFIG_COLUMNS, template_rows),
            "",
            "Low-confidence rows require user confirmation before formal use. Medium-confidence rows should be checked before confirmation.",
        ]) + "\n",
        encoding="utf-8",
    )


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", required=True)
    parser.add_argument("--workspace", required=True)
    parser.add_argument("--source-dir", default="")
    parser.add_argument("--description", default="")
    parser.add_argument("--species-code", default="")
    parser.add_argument("--gene-family", default="")
    parser.add_argument("--primary-reference-species", default="")
    parser.add_argument("--identify-blast-reference-species", default="")
    parser.add_argument("--confirm", action="store_true")
    parser.add_argument("--draft", default="")
    args = parser.parse_args()

    start = time.time()
    input_dir = Path(args.input_dir)
    workspace = Path(args.workspace)
    final_dir = workspace / "final"
    final_dir.mkdir(parents=True, exist_ok=True)
    source_dir = Path(args.source_dir).expanduser() if args.source_dir else input_dir
    draft_path = Path(args.draft).expanduser() if args.draft else final_dir / "config.autofill.draft.xlsx"
    report_path = final_dir / "config.autofill.report.md"
    manifest_path = final_dir / "config_autofill.manifest.json"

    if args.confirm:
        source = draft_path
        if not source.exists():
            raise SystemExit(f"Draft config not found: {source}")
        rows = read_rows(source)
        source_policy = read_analysis_policy(source)
        analysis_family = clean_text(args.gene_family) or source_policy.get("family", "")
        primary_reference_species = clean_text(args.primary_reference_species) or source_policy.get("primary_reference_species", "")
        identify_blast_reference_species = clean_text(args.identify_blast_reference_species) or source_policy.get("identify_blast_reference_species", "")
        archive = backup_if_exists(input_dir / "config.xlsx", input_dir / "uploaded_configs", "config.xlsx")
        input_dir.mkdir(parents=True, exist_ok=True)
        write_workbook(input_dir / "config.xlsx", rows, analysis_family, primary_reference_species, identify_blast_reference_species)
        status = "success"
        message = "Confirmed draft written to input config.xlsx"
    else:
        detected_rows = build_rows(input_dir, source_dir, args.description, args.species_code, args.gene_family)
        existing_rows = read_rows(draft_path) if draft_path.exists() else []
        # Do not preserve a stale Analysis.Family during draft generation.
        # FASTQ-only or genome-resource imports must not silently inherit SPL.
        analysis_family = clean_text(args.gene_family)
        primary_reference_species = clean_text(args.primary_reference_species)
        identify_blast_reference_species = clean_text(args.identify_blast_reference_species)
        rows = merge_rows(existing_rows, detected_rows)
        write_workbook(draft_path, rows, analysis_family, primary_reference_species, identify_blast_reference_species)
        write_report(report_path, rows, source_dir)
        archive = ""
        status = "needs_user_confirmation"
        message = "Draft config generated or merged; user confirmation required before replacing config.xlsx"

    manifest = {
        "manifest_version": "1.0",
        "stage": "config_autofill",
        "status": status,
        "message": message,
        "metrics": {
            "rows": len(rows),
            "new_detected_rows": len(detected_rows) if not args.confirm else 0,
            "existing_draft_rows": len(existing_rows) if not args.confirm else 0,
            "low_confidence_rows": sum(1 for row in rows if row.get("confidence") == "low"),
        },
        "artifacts": {
            "draft_config": {"path": str(draft_path), "type": "xlsx"},
            "report": {"path": str(report_path), "type": "markdown"},
            "confirmed_config": {"path": str(input_dir / "config.xlsx"), "type": "xlsx"} if args.confirm else {},
        },
        "backup": archive,
        "elapsed_sec": round(time.time() - start, 3),
    }
    write_json(manifest_path, manifest)
    print(json.dumps(manifest, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
